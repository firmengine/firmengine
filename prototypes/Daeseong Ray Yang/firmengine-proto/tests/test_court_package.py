"""Output-package tests use synthetic evidence, never held-out court answers."""
from pathlib import Path
import sys
import tempfile
import unittest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.rendering import court_package, render_bundle


def snapshot():
    return {'id': 'synthetic-package', 'revision': 3, 'records': [], 'assertions': [],
            'objects': [], 'documents': [], 'reviews': {}, 'artifacts': []}


def decision(value=True, revision=3):
    return {'value': value, 'state': 'VERIFIED', 'revision': revision, 'reviewer': 'synthetic lawyer'}


class CourtPackageTests(unittest.TestCase):
    def test_existing_evidence_is_not_an_included_court_attachment(self):
        data = snapshot()
        data['documents'] = [{'id': 'doc1', 'name': 'proof.pdf', 'mapping_status': 'VERIFIED'}]
        data['requirements'] = [{'requirement_id': 'req1', 'label': 'Residence proof', 'status': 'SATISFIED',
                                 'matched_document_ids': ['doc1'], 'legal_basis': ['K19']}]
        checks = court_package(data)
        attachment = next(item for item in checks if item['id'] == 'evidence:req1')
        self.assertEqual(attachment['kind'], 'EVIDENCE_ATTACHMENT')
        self.assertEqual(attachment['status'], 'ATTACHMENT_NOT_INCLUDED')
        data['court_attachments'] = [{'requirement_id': 'req1', 'document_id': 'doc1', 'included': True, 'revision': 3}]
        self.assertEqual(next(item for item in court_package(data) if item['id'] == 'evidence:req1')['status'], 'INCLUDED')

    def test_conditional_applications_require_current_explicit_legal_decision(self):
        data = snapshot()
        data['objects'] = [{'id': 'car', 'type': 'vehicle', 'metadata': {'disposal_proposed': True}}]
        data['package_decisions'] = {'exemption_application': decision(revision=2), 'stay_application': decision()}
        with tempfile.TemporaryDirectory() as folder:
            artifacts = render_bundle(data, Path(folder))
        ids = {item['template_id'] for item in artifacts}
        self.assertIn('D5113', ids)
        self.assertNotIn('D5109', ids)
        self.assertNotIn('D5111', ids)
        self.assertNotIn('D5114', ids)

    def test_asset_funded_plan_replaces_income_only_plan_and_keeps_proceeds_schedule(self):
        data = snapshot()
        data['package_decisions'] = {'plan_asset_disposal': decision()}
        with tempfile.TemporaryDirectory() as folder:
            artifacts = render_bundle(data, Path(folder))
            ids = {item['template_id'] for item in artifacts}
            self.assertIn('D5111', ids)
            self.assertNotIn('D5110', ids)
            item = next(item for item in artifacts if item['template_id'] == 'D5111')
            document = Document(Path(folder)/item['name'])
            text = '\n'.join(p.text for p in document.paragraphs)
            self.assertIn('재산처분', text)
            self.assertIn('대체재산', text)

    def test_generated_form_existence_does_not_clear_missing_values_or_manual_signature(self):
        data = snapshot()
        with tempfile.TemporaryDirectory() as folder:
            data['artifacts'] = render_bundle(data, Path(folder))
        checks = court_package(data)
        application = next(item for item in checks if item['id'] == 'application')
        self.assertTrue(application['artifact_present'])
        self.assertEqual(application['status'], 'CONTENT_INCOMPLETE')
        self.assertTrue(application['blocks_filing'])
        self.assertTrue(application['pre_generation_blocker'])
        self.assertTrue(any(item['kind'] == 'MANUAL_SIGNATURE' and item['status'] == 'SIGNATURE_REQUIRED' for item in checks))
        self.assertTrue(all(not item['pre_generation_blocker'] for item in checks if item['kind']=='MANUAL_SIGNATURE'))

    def test_payment_schedule_never_allocates_unreviewed_principal_by_guess(self):
        data = snapshot()
        data['records'] = [{'id': 'r1', 'key': 'creditor_principal', 'entity': 'Synthetic creditor',
                            'value': '1000000', 'state': 'confirmed', 'sources': []}]
        with tempfile.TemporaryDirectory() as folder:
            artifacts = render_bundle(data, Path(folder))
            item = next(item for item in artifacts if item['template_id'] == 'PAYMENT_SCHEDULE')
            doc = Document(Path(folder)/item['name'])
            text = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            self.assertIn('Synthetic creditor', text)
            self.assertIn('[확인 필요]', text)
            # Audit timestamps/hashes may contain these digits; inspect the
            # actual term field rather than every character in the artifact.
            term = next(row.cells[1].text for table in doc.tables for row in table.rows
                        if len(row.cells)==2 and row.cells[0].text=='변제 횟수')
            self.assertEqual(term,'[확인 필요]')

    def test_evidence_group_requires_all_selected_attachments_not_any_uploaded_file(self):
        data = snapshot()
        data['documents'] = [{'id': 'contract', 'document_type': 'lease_contract'}, {'id': 'payment', 'document_type': 'lease_payment_evidence'}]
        data['requirements'] = [{'requirement_id': 'lease-proof', 'label': 'Lease proof', 'status': 'SATISFIED',
            'matched_document_ids': ['contract','payment'], 'alternatives': [{'all_of':['lease_contract','lease_payment_evidence']}]}]
        data['court_attachments'] = [{'requirement_id': 'lease-proof', 'document_id': 'contract', 'included': True, 'revision': 3}]
        self.assertEqual(next(row for row in court_package(data) if row['id']=='evidence:lease-proof')['status'], 'ATTACHMENT_NOT_INCLUDED')
        data['court_attachments'].append({'requirement_id':'lease-proof','document_id':'payment','included':True,'revision':3})
        self.assertEqual(next(row for row in court_package(data) if row['id']=='evidence:lease-proof')['status'], 'INCLUDED')

    def test_canonical_contract_ids_prevent_same_named_creditor_merging(self):
        data = snapshot()
        data['objects'] = [{'id':'a','type':'debt','label':'Same creditor'},{'id':'b','type':'debt','label':'Same creditor'}]
        data['assertions'] = [{'id':'assert-'+oid,'feature_key':'debts[].principal','object_id':oid,'value':amount,
            'state':'VERIFIED','accepted_for_filing':True,'unit':'KRW','valid_at':'2026-10-26','evidence_refs':[]} for oid,amount in [('a','100'),('b','200')]]
        with tempfile.TemporaryDirectory() as folder:
            artifacts = render_bundle(data, Path(folder))
            item = next(row for row in artifacts if row['template_id']=='D5106')
            doc=Document(Path(folder)/item['name'])
            values=[cell.text for table in doc.tables for row in table.rows for cell in row.cells]
            self.assertTrue(any('100' in text for text in values))
            self.assertTrue(any('200' in text for text in values))
            self.assertFalse(any('[불일치·확인 필요]' in text for text in values))

    def test_current_source_backed_allocations_fill_schedule_and_stale_rows_do_not(self):
        data=snapshot()
        data['assertions']=[{'id':'a1','feature_key':'debts[].principal','object_id':'debt1','value':'1000','state':'VERIFIED','accepted_for_filing':True}]
        row={'number':'1','creditor_name':'Synthetic allocation','claim_confirmed':'1000','claim_unconfirmed':'0',
            'monthly_confirmed':'20','monthly_reserved':'0','total_confirmed':'400','total_reserved':'0',
            'case_revision':3,'status':'calculated','source_assertion_ids':['a1'],'basis':'reviewed basis'}
        data['plan_allocations']=[row]
        with tempfile.TemporaryDirectory() as folder:
            artifacts=render_bundle(data,Path(folder))
            doc=Document(Path(folder)/next(item['name'] for item in artifacts if item['template_id']=='PAYMENT_SCHEDULE'))
            body='\n'.join(cell.text for table in doc.tables for tr in table.rows for cell in tr.cells)
            self.assertIn('20 / 0',body)
            self.assertIn('400 / 0',body)
        row['case_revision']=2
        self.assertIn('plan_allocations',next(item for item in court_package(data) if item['id']=='payment_schedule')['missing_features'])

    def test_missing_artifact_alone_does_not_block_generation_when_required_fields_exist(self):
        data=snapshot()
        data['records']=[{'id':key,'key':key,'value':value,'state':'confirmed','sources':[]} for key,value in
            [('client_name','Synthetic'),('resident_id','000000-0000000'),('address','Synthetic address'),
             ('court_name','Synthetic court'),('statement','Synthetic supported statement')]]
        application=next(item for item in court_package(data) if item['id']=='application')
        self.assertEqual(application['status'],'NOT_GENERATED')
        self.assertFalse(application['missing_features'])
        self.assertFalse(application['pre_generation_blocker'])
        self.assertTrue(application['blocks_filing'])


if __name__ == '__main__':
    unittest.main()
