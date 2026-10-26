"""Synthetic artifact tests: no sample.pdf or submitted reference answers used."""
from copy import deepcopy
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.rendering import court_package, render_bundle, review_status


def record(identifier, key, value, *, entity="", period="", state="unreviewed"):
    return {"id": identifier, "key": key, "entity": entity, "label": key,
            "value": value, "unit": "KRW", "period": period, "confidence": 0.9,
            "state": state, "sources": [{"document_id": "synthetic-source", "page": 4,
                                           "quote": "합성 원문 근거"}]}


def text_of(document):
    texts = [paragraph.text for paragraph in document.paragraphs]
    texts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    texts.extend(paragraph.text for section in document.sections for paragraph in section.header.paragraphs)
    return "\n".join(texts)


class RenderingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory()
        cls.session = {"id": "synthetic-render-test", "revision": 3, "template_approved": False,
            "reviews": {}, "records": [
                record("r-name", "client_name", "합성사용자"),
                record("r-2024", "annual_revenue", "12345678901234567890", period="2024", state="confirmed"),
                record("r-2025", "annual_revenue", "40000000", period="2025", state="confirmed"),
                record("r-c1", "creditor_principal", "1000000", entity="합성은행", period="2026-10"),
                record("r-c2", "creditor_interest", None, entity="합성은행", period="2026-10", state="missing"),
                record("r-c3", "creditor_principal", "2000000", entity="합성은행", period="2026-09"),
                record("r-injection", "creditor_address", '=HYPERLINK("https://invalid.test","x")',
                       entity="합성은행", period="2026-10"),
                record("r-address", "address", "충돌주소", state="conflict"),
                record("r-token", "asset_description", "증빙 문자 {{client_name}} 보존", entity="합성재산"),
            ], "calculations": [
                {"key": "monthly_payment", "label": "월 변제액", "value": None,
                 "unit": "KRW", "formula": "income - living - fee", "status": "pending", "sources": []},
                {"key": "debt_principal_total", "label": "검토된 원금 합계", "value": "1000000",
                 "unit": "KRW", "formula": "confirmed principal", "status": "calculated", "sources": ["r-c1"]},
            ]}
        cls.artifacts = render_bundle(cls.session, Path(cls.folder.name))

    @classmethod
    def tearDownClass(cls):
        cls.folder.cleanup()

    def artifact(self, kind):
        return Path(self.folder.name) / next(item["name"] for item in self.artifacts if item["kind"] == kind)

    def test_base_documents_and_applicable_package_annexes_have_actual_form_sections(self):
        self.assertEqual(len(self.artifacts), 10)
        self.assertEqual(sum(item["format"] == "docx" for item in self.artifacts), 9)
        self.assertTrue({'PAYMENT_SCHEDULE','ASSET_VALUATION','EVIDENCE_INDEX'} <= {item['template_id'] for item in self.artifacts})
        for item in self.artifacts:
            self.assertTrue((Path(self.folder.name) / item["name"]).is_file())
        plan = text_of(Document(self.artifact("repayment_plan")))
        self.assertIn("7. 미확정 개인회생채권", plan)
        self.assertIn("개인회생채권 변제예정액표", plan)
        self.assertIn("편집형", plan)
        self.assertIn("미검토 초안", plan)

    def test_missing_and_conflicting_fields_do_not_become_zero_or_default_months(self):
        application = Document(self.artifact("application"))
        rows = {row.cells[0].text: row.cells[1].text for table in application.tables
                for row in table.rows if len(row.cells) == 2}
        self.assertEqual(rows["현주소"], "[불일치·확인 필요]")
        self.assertEqual(rows["변제기간(개월)"], "[확인 필요]")
        self.assertEqual(rows["월 변제예정액"], "[계산 보류]")

    def test_different_periods_are_retained_not_summed(self):
        creditors = text_of(Document(self.artifact("creditors")))
        self.assertIn("2026-10", creditors)
        self.assertIn("2026-09", creditors)
        self.assertIn("1,000,000", creditors)
        self.assertIn("2,000,000", creditors)
        income = Document(self.artifact("income_expense"))
        rows = {row.cells[0].text: row.cells[1].text for table in income.tables
                for row in table.rows if len(row.cells) == 2}
        self.assertEqual(rows["연간 매출(증빙)"], "[기간 선택 필요]")

    def test_provenance_and_literal_evidence_tokens_are_preserved(self):
        assets = text_of(Document(self.artifact("assets")))
        self.assertIn("증빙 문자 {{client_name}} 보존", assets)
        self.assertIn("synthetic-source / 4쪽", assets)
        self.assertIn("SHA-256:", assets)
        self.assertIn("합성 원문 근거", assets)

    def test_workbook_does_not_execute_source_formulas_or_lose_integer_precision(self):
        workbook = load_workbook(self.artifact("workbook"), data_only=False)
        formula_like = [cell for sheet in workbook for row in sheet for cell in row
                        if isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK")]
        self.assertTrue(formula_like)
        self.assertTrue(all(cell.data_type == "s" for cell in formula_like))
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))
        self.assertTrue(any(cell.value == "12345678901234567890" for row in workbook["추출값"] for cell in row))
        workbook.close()

    def test_review_is_bound_to_current_revision_and_template_confirmation(self):
        session = deepcopy(self.session)
        session["reviews"] = {"caseworker": {"name": "직원", "at": "test", "revision": 3},
                              "lawyer": {"name": "변호사", "at": "test", "revision": 3}}
        self.assertIn("미검토", review_status(session))
        session["template_approved"] = True
        self.assertIn("검토 완료본", review_status(session))
        session["revision"] = 4
        self.assertIn("미검토", review_status(session))

    def test_excluded_values_are_not_filled_but_remain_in_evidence_workbook(self):
        session = deepcopy(self.session)
        session['records'] = [record('excluded-name', 'client_name', '제외성명', state='excluded'),
                              record('current-name', 'client_name', '적용성명', state='confirmed'),
                              record('excluded-creditor', 'creditor_principal', '9999999', entity='제외은행', state='excluded')]
        with tempfile.TemporaryDirectory() as directory:
            files = render_bundle(session, Path(directory))
            for artifact in files:
                if artifact['format'] == 'docx':
                    text = text_of(Document(Path(directory) / artifact['name']))
                    self.assertNotIn('제외성명', text)
                    self.assertNotIn('제외은행', text)
            workbook_path = Path(directory) / next(item['name'] for item in files if item['format'] == 'xlsx')
            workbook = load_workbook(workbook_path)
            values = [cell.value for row in workbook['추출값'] for cell in row]
            self.assertIn('제외성명', values)
            self.assertIn('검토 제외', values)
            workbook.close()

    def test_digit_only_identifiers_keep_leading_zeroes_in_form_body(self):
        session = deepcopy(self.session)
        session['records'] = [
            record('person-id', 'resident_id', '0001011234567'),
            record('phone', 'phone', '01000112233'),
            record('business-id', 'business_registration', '0012345678'),
            record('bank', 'bank_balance', '1234567'),
        ]
        with tempfile.TemporaryDirectory() as directory:
            artifacts = render_bundle(session, Path(directory))
            paths = {item['kind']: Path(directory) / item['name'] for item in artifacts}
            application = Document(paths['application'])
            rows = {row.cells[0].text: row.cells[1].text for table in application.tables
                    for row in table.rows if len(row.cells) == 2}
            self.assertEqual(rows['주민등록번호'], '0001011234567')
            self.assertEqual(rows['휴대전화'], '01000112233')
            self.assertEqual(rows['사업자등록번호'], '0012345678')
            assets = Document(paths['assets'])
            asset_rows = {row.cells[0].text: row.cells[1].text for table in assets.tables
                          for row in table.rows if len(row.cells) == 2}
            self.assertEqual(asset_rows['예금'], '1,234,567')

    def test_undated_creditor_metadata_fills_matching_snapshot_rows_without_merging_contracts(self):
        session = deepcopy(self.session)
        session['records'] = [
            record('p1', 'creditor_principal', '1000', entity='합성 은행 / 계약 A', period='2026-09'),
            record('i1', 'creditor_interest', '10', entity='합성은행/계약A', period='2026-09'),
            record('p2', 'creditor_principal', '2000', entity='합성 은행 / 계약 A', period='2026-10'),
            record('i2', 'creditor_interest', '20', entity='합성은행/계약A', period='2026-10'),
            record('a', 'creditor_address', '합성 주소 A', entity='합성은행/계약A'),
            record('p3', 'creditor_principal', '3000', entity='합성은행/계약B', period='2026-10'),
            record('i3', 'creditor_interest', '0', entity='합성은행/계약B', period='2026-10'),
            record('b', 'creditor_address', '합성 주소 B', entity='합성은행/계약B'),
        ]
        with tempfile.TemporaryDirectory() as directory:
            artifacts = render_bundle(session, Path(directory))
            path = Path(directory) / next(item['name'] for item in artifacts if item['kind'] == 'workbook')
            workbook = load_workbook(path)
            rows = list(workbook['채권자'].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 3)
            self.assertEqual([(row[1], row[2], row[3], row[4]) for row in rows], [
                ('2026-09', '1,000', '10', '합성 주소 A'),
                ('2026-10', '2,000', '20', '합성 주소 A'),
                ('2026-10', '3,000', '0', '합성 주소 B'),
            ])
            workbook.close()

    def test_source_backed_draft_fields_and_selected_creditors_fill_body_without_overriding_confirmed_values(self):
        session = deepcopy(self.session)
        session['records'] = [record('confirmed-name', 'client_name', '확정 합성인', state='confirmed')]
        support = {'status': 'proposed', 'source_pages': [{'document_id': 'synthetic-source', 'page': 4}],
                   'reason': '합성 증빙에 근거한 검토 전 제안', 'kind': 'summary'}
        session['draft_fields'] = {'client_name': {**support, 'value': '제안 합성인'},
                                   'education': {**support, 'value': '합성고등학교 졸업'}}
        session['draft_creditors'] = [{**support, 'name': '선정 합성은행', 'contract_id': '00123',
            'period': '2026-10-01', 'principal': '1234567', 'interest': '3456', 'address': '합성 주소'}]
        session['draft_calculations'] = [{'key': 'monthly_payment', 'value': '456789', 'status': 'proposed',
                                          'unit': 'KRW', 'formula': '합성 입력으로 계산'}]
        with tempfile.TemporaryDirectory() as directory:
            files = render_bundle(session, Path(directory))
            paths = {item['kind']: Path(directory) / item['name'] for item in files}
            statement = Document(paths['statement'])
            fields = {row.cells[0].text: row.cells[1].text for table in statement.tables
                      for row in table.rows if len(row.cells) == 2}
            self.assertEqual(fields['최종 학력'], '합성고등학교 졸업')
            application = Document(paths['application'])
            fields = {row.cells[0].text: row.cells[1].text for table in application.tables
                      for row in table.rows if len(row.cells) == 2}
            applicant = next(row.cells[1].text for table in application.tables for row in table.rows
                             if len(row.cells) == 2 and row.cells[0].text == '성명')
            self.assertEqual(applicant, '확정 합성인')
            self.assertEqual(fields['월 변제예정액'], '456,789')
            creditors = Document(paths['creditors'])
            body = next(table for table in creditors.tables if len(table.columns) >= 4)
            self.assertIn('선정 합성은행', body.rows[1].cells[0].text)
            self.assertIn('1,234,567', ' '.join(cell.text for cell in body.rows[1].cells))
            self.assertIn('미검토', review_status(session))

    def test_public_form_detail_tables_contain_evidence_in_body(self):
        session = deepcopy(self.session)
        session['records'] = [
            record('p', 'creditor_principal', '1230000', entity='합성채권자', period='2026-10-01'),
            record('i', 'creditor_interest', '23000', entity='합성채권자', period='2026-10-01'),
            record('cause', 'creditor_cause', '합성 대여계약', entity='합성채권자'),
            record('basis', 'creditor_principal_basis', '합성 부채확인서 잔액', entity='합성채권자'),
            record('v', 'asset_value', '45000', entity='합성예금', period='2026-10-01'),
            record('account', 'asset_account_number', '00123400001', entity='합성예금'),
            record('relation', 'family_relationship', '자녀', entity='합성가족'),
            record('occupation', 'family_occupation', '학생', entity='합성가족'),
        ]
        with tempfile.TemporaryDirectory() as directory:
            files = render_bundle(session, Path(directory))
            paths = {item['kind']: Path(directory) / item['name'] for item in files}
            creditors = Document(paths['creditors'])
            body = next(table for table in creditors.tables if len(table.columns) >= 4)
            details = '\n'.join(cell.text for row in body.rows for cell in row.cells)
            self.assertIn('합성 부채확인서 잔액', details)
            self.assertIn('합성 대여계약', details)
            assets = Document(paths['assets'])
            body_details = '\n'.join(cell.text for table in assets.tables if len(table.columns) >= 3
                                     for row in table.rows for cell in row.cells)
            self.assertIn('00123400001', body_details)
            income = Document(paths['income_expense'])
            family = next(table for table in income.tables if '관계' in [cell.text for cell in table.rows[0].cells])
            self.assertIn('합성가족', '\n'.join(cell.text for row in family.rows for cell in row.cells))
            plan = Document(paths['repayment_plan'])
            boilerplate = '\n'.join(paragraph.text for paragraph in plan.paragraphs)
            self.assertIn('제625조', boilerplate)

    def test_verified_value_requires_filing_adoption_and_cannot_fall_back_to_old_legacy_value(self):
        session = deepcopy(self.session)
        session['records'] = [record(key, key, value, state='confirmed') for key, value in [
            ('client_name', '이전 채택값'), ('resident_id', '000000-0000000'),
            ('address', '합성 주소'), ('court_name', '합성 법원'), ('statement', '합성 진술')]]
        assertion = {'id': 'verified-name', 'feature_key': 'person.name', 'object_id': 'person-primary',
                     'value': '미채택 검증성명', 'state': 'VERIFIED', 'accepted_for_filing': True,
                     'evidence_refs': [{'document_id': 'synthetic-source', 'page': 4, 'excerpt': '합성 원문'}]}
        session['assertions'] = [assertion]
        check = next(item for item in court_package(session) if item['id'] == 'application')
        self.assertNotIn('client_name', check['missing_features'])
        self.assertFalse(check['pre_generation_blocker'])
        assertion['accepted_for_filing'] = False
        check = next(item for item in court_package(session) if item['id'] == 'application')
        self.assertIn('client_name', check['missing_features'])
        self.assertTrue(check['pre_generation_blocker'])
        with tempfile.TemporaryDirectory() as folder:
            artifacts = render_bundle(session, Path(folder))
            application = Document(Path(folder) / next(item['name'] for item in artifacts if item['kind'] == 'application'))
            name = next(row.cells[1].text for table in application.tables for row in table.rows
                        if len(row.cells) == 2 and row.cells[0].text == '성명')
            self.assertEqual(name, '[확인 필요]')
            self.assertNotIn('이전 채택값', name)
            self.assertNotIn('미채택 검증성명', name)
            workbook = load_workbook(Path(folder) / next(item['name'] for item in artifacts if item['kind'] == 'workbook'))
            try:
                values = [cell.value for row in workbook['사건주장근거'] for cell in row]
                self.assertIn('미채택 검증성명', values)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
