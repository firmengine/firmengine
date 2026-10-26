"""Rebuild editable template assets using this directory's public reference forms.

Development-time utility only. No external project is needed for rebuilding.
Changing a template changes its hash and requires a new review of the mapping.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
VERSION = "public-form-editable-reconstruction-v2"


def token(key):
    return "{{" + key + "}}"


def start(code, title):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.9)
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(10 if name == "Normal" else 12)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    section.header.paragraphs[0].text = f"{code}  |  {token('review_status')}  |  버전 {token('revision')}"
    section.header.paragraphs[0].runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.text = "공식 양식 항목의 편집형 재구성 · 서명·동의 별도 확인  |  "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.runs[0].font.size = Pt(8)
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("공식 원본의 구역·항목을 DOCX로 재구성한 초안입니다. 원본과 동일한 페이지 배치는 아니며, 제공 양식과의 대응 검토가 필요합니다.")
    return doc


def table(doc, headers, rows):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.style = "Table Grid"
    for cell, value in zip(tab.rows[0].cells, headers):
        cell.text = value
    props = tab.rows[0]._tr.get_or_add_trPr()
    props.append(OxmlElement("w:tblHeader"))
    for values in rows:
        for cell, value in zip(tab.add_row().cells, values):
            cell.text = value
    doc.add_paragraph("")
    return tab


def fields(doc, rows):
    return table(doc, ["항목", "기재 내용"], [[label, token(key)] for label, key in rows])


def section(doc, title, rows):
    doc.add_heading(title, level=1)
    fields(doc, rows)


def sign(doc):
    doc.add_paragraph("작성일: " + token("application_date"))
    doc.add_paragraph("신청인 / 채무자: " + token("client_name") + "    (서명 또는 날인: __________________)")
    doc.add_paragraph(token("court_name") + " 귀중")


def application():
    doc = start("D5100", "개인회생절차 개시신청서")
    section(doc, "신청인", [("성명", "client_name"), ("주민등록번호", "resident_id"),
        ("주민등록상 주소", "registered_address"), ("현주소", "address"),
        ("송달장소", "service_address"), ("휴대전화", "phone")])
    section(doc, "대리인", [("성명", "lawyer_name"), ("사무실 주소", "lawyer_address"),
        ("전화번호", "lawyer_phone"), ("이메일", "lawyer_email")])
    section(doc, "관련 사건", [("주채무자·보증인·연대채무자·배우자의 관련 사건", "related_cases")])
    doc.add_heading("신청취지", level=1)
    doc.add_paragraph("「신청인에 대하여 개인회생절차를 개시한다.」라는 결정을 구합니다.")
    section(doc, "신청이유 및 수입 상황", [("사업장명", "business_name"), ("업종", "business_type"),
        ("사업자등록번호", "business_registration"), ("사건 경위", "statement")])
    doc.add_paragraph("정기·반복 수입, 지급곤란 및 기각사유 해당 여부는 담당자·변호사가 검토합니다. 해당 선택란을 자동으로 확인하지 않습니다.")
    doc.add_paragraph("□ 신청인은 별지 개인회생채권자목록 기재와 같은 채무를 부담하고 있으나, 수입 및 재산이 별지 수입 및 지출에 관한 목록과 재산목록에 기재된 바와 같으므로, 파산의 원인사실이 발생하였습니다(파산의 원인사실이 생길 염려가 있습니다).")
    doc.add_paragraph("□ 신청인은 부동산임대소득·사업소득·농업소득·임업소득 그 밖에 이와 유사한 수입을 장래에 계속적으로 또는 반복하여 얻을 것으로 예상되고, 또한 채무자 회생 및 파산에 관한 법률 제595조에 해당하는 개시신청 기각사유는 없습니다(영업소득자의 경우).")
    section(doc, "변제예정 및 적립금 반환계좌", [("변제기간(개월)", "repayment_months"),
        ("월 변제예정액", "monthly_payment"), ("첫 납입일", "repayment_start_date"),
        ("매월 납입일", "repayment_day"), ("반환은행", "refund_bank"), ("반환계좌", "refund_account")])
    doc.add_paragraph("신청인은 각 회생채권자에 대한 채무 전액의 변제가 곤란하므로, 그 일부를 분할하여 지급할 계획입니다. 현시점에서 계획하고 있는 변제예정액은 " + token('repayment_months') + "개월간 월 " + token('monthly_payment') + "원씩입니다. 개시결정이 내려지는 경우 " + token('repayment_start_date') + "을 제1회로 하여, 이후 매월 " + token('repayment_day') + "일에 개시결정 시 통지되는 개인회생위원의 은행계좌에 입금하는 내용으로 검토합니다.")
    doc.add_paragraph("이 사건 개인회생절차에서 적립금을 반환받을 신청인의 예금계좌는 " + token('refund_bank') + " " + token('refund_account') + "입니다. 계좌가 변경되거나 사용할 수 없게 된 경우 사건담당 회생위원에게 통장사본을 첨부하여 신고하는 내용을 확인합니다.")
    doc.add_heading("첨부서류", level=1)
    for label in ["개인회생채권자목록", "재산목록", "수입 및 지출에 관한 목록", "진술서", "예금계좌 사본", "위임장 및 비용납부 관련 자료"]:
        doc.add_paragraph("☐ " + label + " (실제 첨부 여부 확인)")
    sign(doc)
    doc.add_heading("휴대전화를 통한 정보수신 신청", level=1)
    doc.add_paragraph("☐ 신청 여부 별도 확인    휴대전화: " + token("phone"))
    doc.add_paragraph("신청인: __________________ (서명 또는 날인)")
    return doc


def creditors():
    doc = start("D5106", "개인회생채권자목록")
    section(doc, "작성 기준", [("채권현재액 산정기준일", "debt_as_of"), ("목록 작성일", "application_date"),
        ("원금 합계", "total_principal"), ("이자 합계", "total_interest"),
        ("담보부 회생채권액 합계", "secured_debt_total"), ("무담보 회생채권액 합계", "unsecured_debt_total")])
    doc.add_heading("채권자별 내역", level=1)
    table(doc, ["채권번호·채권자 / 기준일", "채권의 원인·내용", "주소 및 연락처", "채권현재액 및 산정근거"], [[
        "{{creditor.number}}. {{creditor.entity}}\n기준일: {{creditor.period}}\n계약·계좌: {{creditor.contract_id}}",
        "발생원인: {{creditor.cause}}\n발생일: {{creditor.loan_date}}\n최초원금: {{creditor.original_principal}}\n내용: {{creditor.content}}\n유형: {{creditor.claim_type}}",
        "주소: {{creditor.address}}\n전화: {{creditor.phone}}\n팩스: {{creditor.fax}}",
        "원금: {{creditor.principal}} 원\n원금 산정근거: {{creditor.principal_basis}}\n이자: {{creditor.interest}} 원\n이자 산정근거: {{creditor.interest_basis}}"]])
    doc.add_paragraph("동일 채권자의 서로 다른 기간·계약은 자동 합산하지 않습니다. 채권 순위·동일성·담보 분류는 별도 검토합니다.")
    section(doc, "부속서류 1. 별제권부채권 및 이에 준하는 채권", [("해당 여부·목적물·예상 변제액·담보부족액", "secured_claim_details")])
    table(doc, ["채권번호·채권자", "원금 / 이자", "예상 변제액 / 부족액 / 담보부 채권액", "별제권의 내용 및 목적물"], [[
        "{{secured.number}}. {{secured.entity}}", "{{secured.principal}} / {{secured.interest}}",
        "③ {{secured.secured_recovery}}\n④ {{secured.shortfall}}\n⑤ {{secured.secured_amount}}", "{{secured.collateral}}"]])
    section(doc, "부속서류 2. 다툼이 있거나 예상되는 채권", [("채권번호·주장액·다툼 없는 부분·차이·소송경과", "disputed_claim_details")])
    table(doc, ["채권번호·채권자", "채권자 주장액 / 다툼 없는 부분", "다툼의 원인", "소송제기 여부 및 진행경과"], [[
        "{{disputed.number}}. {{disputed.entity}}", "{{disputed.disputed_amount}} / {{disputed.undisputed_amount}}",
        "{{disputed.dispute_reason}}", "{{disputed.litigation}}"]])
    section(doc, "부속서류 3. 전부명령", [("채권자·채권내용·전부명령 내역", "assignment_order_details")])
    table(doc, ["채권번호·채권자", "채권의 내용", "전부명령의 내역"], [[
        "{{assignment.number}}. {{assignment.entity}}", "{{assignment.content}}", "{{assignment.assignment_order}}"]])
    section(doc, "부속서류 4. 기타", [("보증채무·제3자 담보 등 기타 내역", "other_claim_details")])
    sign(doc)
    return doc


def assets():
    doc = start("D5101", "재산목록")
    fields(doc, [("현금", "cash"), ("예금", "bank_balance"), ("보험 해약환급금", "insurance_surrender"),
        ("자동차", "vehicle_value"), ("임차보증금 반환채권", "housing_deposit"), ("부동산", "real_estate_value"),
        ("대여금 채권", "loan_receivables"), ("매출금 채권", "trade_receivables"), ("예상퇴직금", "retirement_value"),
        ("사업용 설비·재고품·비품", "business_assets_value"),
        ("기타 재산", "other_assets"), ("재산 합계", "assets_total"), ("면제재산 신청금액", "exempt_value"),
        ("면제재산 결정신청 내용", "exempt_details"),
        ("청산가치", "liquidation_value")])
    doc.add_heading("재산별 명세", level=1)
    table(doc, ["명칭 / 기준일", "금액 또는 시가 (원)", "압류·가압류 유무", "비고 및 세부명세"], [[
        "{{asset.entity}}\n종류: {{asset.kind}}\n기준일: {{asset.period}}", "{{asset.value}}",
        "{{asset.seizure}}", "{{asset.description}}\n기관·상대방: {{asset.institution}}\n계좌·증권·등록번호: {{asset.account_number}}\n소재지: {{asset.address}}\n취득시기: {{asset.acquisition_date}}\n소유지분: {{asset.ownership_share}}\n담보내역: {{asset.collateral}}\n가액 산정근거: {{asset.valuation_basis}}\n반환 예상금액: {{asset.return_amount}}"]])
    doc.add_paragraph("증빙의 원금·가액을 청산가치로 자동 인정하지 않습니다. 평가기준일, 담보, 면제·공제 및 중복 계상을 확인해야 합니다.")
    sign(doc)
    return doc


def income():
    doc = start("D5103", "수입 및 지출에 관한 목록")
    section(doc, "I. 현재의 수입 목록 — 영업소득", [("상호 / 사업장명", "business_name"), ("업종", "business_type"),
        ("과세유형", "tax_type"), ("사업자등록번호", "business_registration"), ("사업장 주소", "business_address"),
        ("개업일", "opening_date"), ("소득 귀속연도", "income_year"), ("연간 매출(증빙)", "annual_revenue"),
        ("연간 경비(증빙)", "annual_expenses"), ("연간 사업소득(증빙)", "annual_business_income"),
        ("월평균 소득(담당자 인정값)", "approved_monthly_income"), ("검토된 사업소득 계산", "monthly_business_income")])
    doc.add_heading("월별 매출 및 사업경비 근거", level=2)
    table(doc, ["대상 월 / 자료기간", "매출", "사업경비", "자료상 차액 (인정 전)"], [["{{month.entity}}\n{{month.period}}", "{{month.revenue}}", "{{month.expense}}", "{{month.net_income}}"]])
    section(doc, "II. 변제계획 수행 시의 예상 지출 목록", [("인정 가구원수", "household_size"),
        ("인정 생계비", "recognized_living_cost"), ("주거비", "housing_cost"), ("의료비", "medical_cost"),
        ("교육비", "education_cost"), ("추가 생계비 사유", "additional_living_reason")])
    doc.add_paragraph("생계비 비율과 추가 인정 여부의 선택란은 자동 선택하지 않습니다.")
    section(doc, "III. 가족관계", [("관계·동거·직업·월수입·재산·부양 여부", "family_details")])
    table(doc, ["관계", "성명", "연령", "직업", "월수입", "재산가액", "동거 여부", "부양 여부"], [[
        "{{family.relationship}}", "{{family.entity}}", "{{family.age}}", "{{family.occupation}}",
        "{{family.monthly_income}}", "{{family.assets}}", "{{family.cohabitation}}", "{{family.dependent}}"]])
    sign(doc)
    return doc


def statement():
    doc = start("D5105", "진술서")
    section(doc, "I. 경력", [("최종 학력", "education"), ("사업장명", "business_name"),
        ("업종", "business_type"), ("개업일", "opening_date"), ("과거 경력", "career_history"),
        ("과거 결혼·이혼 경력", "marriage_history")])
    table(doc, ["기간", "직장·사업장명", "업종", "직위·자영 여부"], [[
        "{{career.period}}", "{{career.employer}}", "{{career.business_type}}", "{{career.position}}"]])
    section(doc, "II. 현재 주거 상황", [("주소", "address"), ("거주 시작일", "residence_start"),
        ("거주 형태", "housing_type"), ("임차보증금", "housing_deposit"), ("월 임대료", "housing_cost"),
        ("임대료 연체액", "housing_arrears"), ("임차인 성명", "housing_tenant"),
        ("주택 소유자", "housing_owner"), ("신청인과의 관계", "housing_owner_relationship")])
    section(doc, "III. 부채 상황", [("소송·지급명령·전부명령·압류 등", "enforcement_history")])
    table(doc, ["내역", "채권자", "관할법원", "사건번호"], [[
        "{{enforcement.kind}}", "{{enforcement.creditor}}", "{{enforcement.court}}", "{{enforcement.case_number}}"]])
    fields(doc, [("개인회생절차에 이르게 된 사정 (진술에 명시된 항목)", "debt_cause_categories")])
    doc.add_heading("많은 채무를 부담하게 된 사정 및 신청에 이르게 된 사정", level=2)
    doc.add_paragraph(token("statement"))
    doc.add_paragraph("위 문장은 연결된 진술·증빙 범위에서 작성된 초안입니다. 확인되지 않은 생활·사업 경위는 담당자가 보완합니다.")
    section(doc, "IV. 과거 면책절차 등의 이용 상황", [("절차·법원 또는 기관·신청시기·진행상황", "prior_proceedings")])
    table(doc, ["절차", "법원 또는 기관", "신청 시기", "현재까지 진행상황"], [[
        "{{prior.kind}}", "{{prior.court}}", "{{prior.period}}", "{{prior.status}}"]])
    sign(doc)
    return doc


def plan():
    doc = start("D5110", "변제계획안")
    doc.add_paragraph("원본 유형: 가용소득만으로 변제하는 경우. 해당 유형의 적합성은 변호사가 확인합니다.")
    section(doc, "변제계획안 제출서", [("사건번호", "case_number"), ("채무자", "client_name"), ("대리인", "lawyer_name")])
    doc.add_paragraph("채무자는 별지와 같이 변제계획안을 작성하여 제출하니 인가하여 주시기 바랍니다.")
    section(doc, "1. 변제기간", [("시작일", "repayment_start_date"), ("종료일", "repayment_end_date"), ("변제 개월 수", "repayment_months")])
    section(doc, "2. 변제에 제공되는 소득 또는 재산", [("운영 사업장", "business_name"),
        ("월평균 수입(인정값)", "approved_monthly_income"), ("채무자 및 피부양자 수", "household_size"),
        ("인정 생계비", "recognized_living_cost"), ("월평균 가용소득", "monthly_disposable_income"),
        ("월 회생위원 보수", "monthly_trustee_fee"), ("월 실제 변제예정액", "monthly_payment"),
        ("총 변제예정액", "total_repayment"), ("재산 처분 여부·내역", "asset_disposal_plan")])
    fields(doc, [("기준 중위소득에 따른 생계비 (적용 연도·가구수 별도 확인)", "standard_living_cost")])
    table(doc, ["① 월평균 수입", "② 월평균 생계비", "③ 가용소득 (①-②)", "④ 월 회생위원 보수", "⑤ 실제 가용소득 (③-④)", "⑥ 변제횟수", "⑦ 총 실제 가용소득"], [[
        token('approved_monthly_income'), token('recognized_living_cost'), token('monthly_disposable_income'),
        token('monthly_trustee_fee'), token('monthly_payment'), token('repayment_months'), token('total_repayment')]])
    section(doc, "3. 개인회생재단채권에 대한 변제", [("회생위원 보수·비용 및 기타 재단채권", "estate_claim_plan")])
    doc.add_paragraph("[해당 사항 있음 □ / 해당 사항 없음 □] 회생위원 보수·비용과 기타 재단채권의 내용, 금액 및 변제기를 확인합니다. 정형 변제방법: 변제계획 인가일 직후 원리금 전액을 일반 개인회생채권보다 우선하여 변제한다. 적용 여부는 검토 전입니다.")
    table(doc, ["채권자", "채권현재액", "채권발생원인", "변제기"], [[
        '{{estate.entity}}', '{{estate.principal}}', '{{estate.cause}}', '{{estate.due_date}}']])
    section(doc, "4. 일반의 우선권 있는 개인회생채권", [("채권내용·우선권 근거·변제방법", "priority_claim_plan")])
    table(doc, ["채권자", "채권현재액", "채권발생원인 (우선권 근거)", "변제기"], [[
        '{{priority.entity}}', '{{priority.principal}}', '{{priority.priority_basis}}', '{{priority.due_date}}']])
    doc.add_paragraph("[해당 사항 있음 □ / 해당 사항 없음 □] 정형 변제방법: 변제계획 인가일 직후 최초 도래하는 변제기일에 원리금 전액을 우선하여 변제한다. 남은 채권이 있을 경우에는 일반 개인회생채권의 매 변제기일에 우선하여 변제한다. 적용 여부는 검토 전입니다.")
    section(doc, "5. 별제권부 채권 및 이에 준하는 채권", [("채권내용·별제권 목적물·변제방법", "secured_claim_details")])
    table(doc, ["채권번호·채권자", "원금 / 이자", "③ 예상 변제액 / ④ 담보부족액", "별제권 등의 내용 및 목적물"], [[
        '{{secured.number}}. {{secured.entity}}', '{{secured.principal}} / {{secured.interest}}',
        '{{secured.secured_recovery}} / {{secured.shortfall}}', '{{secured.collateral}}']])
    doc.add_paragraph("[해당 사항 있음 □ / 해당 사항 없음 □] 별제권 행사 등으로 변제가 예상되는 채권액은 별제권 행사 등에 의한 방법으로 변제하고 이 변제계획상의 가용소득이나 재산처분에 의한 변제대상에서 제외한다. 별제권 행사 등으로도 변제받을 수 없을 채권액은 미확정채권으로 보아 유보하였다가 아래 7항 기재와 같은 방법으로 변제한다. 적용 범위·금액은 검토 전입니다.")
    section(doc, "6. 일반 개인회생채권에 대한 변제", [("월 변제예정액", "monthly_payment"),
        ("총 변제예정액", "total_repayment"), ("변제일·회차·방법", "payment_schedule_details")])
    fields(doc, [("안분 기준 (원금·원리금·특수채권 포함 여부 검토)", "allocation_basis")])
    doc.add_paragraph("각 일반 개인회생채권의 " + token('allocation_basis') + "을 기준으로 월 평균가용소득에서 월 회생위원 보수를 차감한 월 실제 가용소득을 안분하여 산출한 금액을 변제하는 안입니다. 월 변제예정(유보)액은 " + token('monthly_payment') + "원, 총 변제예정(유보)액은 " + token('total_repayment') + "원이며 구체적 산정내역은 아래 변제예정액표와 같습니다.")
    doc.add_paragraph("기간 및 횟수: " + token('repayment_start_date') + "부터 " + token('repayment_end_date') + "까지 " + token('repayment_months') + "개월간. 변제월 및 변제일: 매월 " + token('repayment_day') + "일 (인가 전 적립금의 일시변제·별도 주기 여부는 개별 확인).")
    section(doc, "7. 미확정 개인회생채권에 대한 조치", [("유보·확정 후 변제방법", "unconfirmed_claim_plan")])
    doc.add_paragraph("[해당 사항 있음 □ / 해당 사항 없음 □] 미확정 개인회생채권에 대하여는 변제를 유보하고, 별지 개인회생채권 변제예정액표에 기재한 금액을 당해 채권이 확정될 때까지 유보하여 둔다. 유보금액도 다른 채권의 변제금과 함께 아래 8항의 계좌에 입금한다.")
    doc.add_paragraph("미확정 개인회생채권의 전액이 확정된 경우에는 그 확정 직후 유보비율을 변제비율로 적용하여 변제를 개시하며, 이미 분할 변제기가 도래한 유보액은 일시에 변제한다. 일부 부존재로 확정되거나 변제기간 종료까지 미확정인 경우의 안분 기준 및 미사용 유보액 처리는 원본 양식과 적용 채권의 확정내용을 대조하여 보완한다.")
    section(doc, "8. 변제금원의 회생위원에 대한 임치 및 지급", [("회생위원 예금계좌·지급방법", "trustee_payment_details")])
    doc.add_paragraph("회생위원 계좌는 확인된 경우에만 기재합니다. 임의의 계좌를 생성하지 않습니다.")
    doc.add_paragraph("채무자는 변제하여야 할 금액을 개시결정 시 통지되는 개인회생위원의 예금계좌에 순차 임치하고, 개인회생채권자는 법원에 예금계좌를 신고하여 회생위원으로부터 변제액을 송금받는 방법으로 지급받는다. 회생위원은 계좌번호를 신고하지 않은 개인회생채권자에 대하여는 변제액을 적립하였다가 이를 연 1회 개인회생사건이 계속되어 있는 지방법원에 공탁하여 지급할 수 있다.")
    section(doc, "9. 면책의 범위 및 효력발생시기", [("해당 법률·예외·개별 조건 검토", "discharge_clause_review")])
    doc.add_paragraph("채무자가 개인회생채권에 대하여 이 변제계획에 따라 변제를 완료하고 면책신청을 하여 면책결정이 확정되었을 경우에는, 이 변제계획에 따라 변제한 것을 제외하고 개인회생채권자에 대한 채무에 관하여 그 책임이 면제된다. 단, 채무자 회생 및 파산에 관한 법률 제625조 제2항 단서 각호 소정의 채무에 관하여는 그러하지 아니하다. (공개 원본의 정형문구이며, 사건 적용·현행성 검토 전)")
    section(doc, "10. 기타 사항", [("별도 기재사항", "plan_other_details")])
    doc.add_heading("개인회생채권 변제예정액표", level=1)
    table(doc, ["채권번호·채권자 / 기준기간", "채권액 (원금)", "유형·확정성", "월 변제예정(유보)액", "총 변제예정(유보)액"], [[
        "{{creditor.number}}. {{creditor.entity}}\n{{creditor.period}}", "{{creditor.principal}}", "{{creditor.claim_type}}", "{{creditor.monthly_payment}}", "{{creditor.total_payment}}"]])
    doc.add_paragraph("채권자별 배분·우선권·미확정채권 유보는 별도 확인이 필요합니다. 총액만으로 임의 안분하지 않습니다.")
    fields(doc, [("변제율(%)", "repayment_rate"), ("청산가치", "liquidation_value"),
        ("총 변제예정액", "total_repayment"), ("현재가치(동일 기준일 검토)", "present_value")])
    sign(doc)
    return doc


def build():
    ROOT.mkdir(exist_ok=True)
    (ROOT / "originals").mkdir(exist_ok=True)
    # The copied, hash-pinned public originals and catalog are the authoring
    # source. Rebuilding remains possible after the external references are removed.
    existing = json.loads((ROOT / "catalog.json").read_text(encoding="utf-8"))
    sources = {(form["code"], Path(source["path"]).suffix.lstrip(".")): source
               for form in existing["forms"] for source in form["originals"]}
    constructors = [
        ("D5100", "application", "개인회생절차 개시신청서", application),
        ("D5106", "creditors", "개인회생채권자목록", creditors),
        ("D5101", "assets", "재산목록", assets),
        ("D5103", "income_expense", "수입 및 지출에 관한 목록", income),
        ("D5105", "statement", "진술서", statement),
        ("D5110", "repayment_plan", "변제계획안", plan),
    ]
    import re
    forms = []
    for code, kind, title, constructor in constructors:
        original_rows = []
        for extension in ["pdf", "hwp"]:
            source = sources[(code, extension)]
            data = (ROOT / source["path"]).read_bytes()
            digest = hashlib.sha256(data).hexdigest()
            if digest != source["sha256"]:
                raise ValueError("Original source hash mismatch: " + code)
            relative = f"originals/{code}.{extension}"
            (ROOT / relative).write_bytes(data)
            original_rows.append({"path": relative, "title": source["title"], "sha256": digest,
                "url": source["url"], "retrieved_at": source["retrieved_at"]})
        doc = constructor()
        filename = code + ".docx"
        doc.save(ROOT / filename)
        texts = [p.text for p in doc.paragraphs]
        texts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        tokens = sorted(set(re.findall(r"\{\{([a-zA-Z0-9_.-]+)\}\}", "\n".join(texts))))
        record_keys = [key for key in tokens if "." not in key]
        aliases = {name: 'creditor' for name in ('secured', 'disputed', 'assignment', 'priority', 'estate')}
        aliases['prior'] = 'prior_procedure'
        for key in tokens:
            if '.' not in key:
                continue
            prefix, suffix = key.split('.', 1)
            if suffix in ('entity', 'period', 'monthly_payment', 'total_payment', 'net_income'):
                continue
            mapped = ('monthly_' + suffix if prefix == 'month' else aliases.get(prefix, prefix) + '_' + suffix)
            record_keys.append(mapped)
        # Provenance includes the extraction candidates used for recognized income.
        if code in ["D5103", "D5110"]:
            record_keys.extend(["monthly_income", "monthly_revenue", "monthly_expense"])
        calc_keys = ["debt_principal_total", "debt_interest_total", "monthly_business_income",
                     "monthly_disposable_income", "monthly_payment", "total_repayment", "repayment_rate"]
        forms.append({"code": code, "kind": kind, "title": title, "template_version": VERSION,
            "template": {"path": filename, "sha256": hashlib.sha256((ROOT / filename).read_bytes()).hexdigest()},
            "originals": original_rows, "record_keys": sorted(set(record_keys)),
            "calculation_keys": calc_keys, "mapped_tokens": tokens,
            "approval_status": "reviewer_pending", "reconstruction": True,
            "limitations": "Editable reconstruction of official section/field structure, not a pixel-exact HWP conversion. Missing and unsupported fields remain explicit. Court/date applicability and field mapping require expert approval."})
    catalog = {"version": VERSION, "created_on": "2026-10-26", "jurisdiction": "national_common_reference",
        "source_snapshot": "2026-10-03", "approval_status": "reviewer_pending", "forms": forms}
    from package_templates import build_additional
    catalog = build_additional(catalog)
    (ROOT / "catalog.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Built {len(catalog['forms'])} editable templates with hash-verified public originals.")


if __name__ == "__main__":
    build()
