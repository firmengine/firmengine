"""Synthetic regression checks for strict primary-body evaluation. No API calls."""
from pathlib import Path
import tempfile
import unittest
from docx import Document
from openpyxl import Workbook

import evaluate_sample as scalar
import submission_quality as quality


class SubmissionQualityTests(unittest.TestCase):
    def test_appendix_value_never_fills_an_empty_primary_field(self):
        quality.PRIVATE.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=quality.PRIVATE, prefix="body_test_") as directory:
            path = Path(directory) / "draft.docx"
            doc = Document()
            doc.add_heading("재산목록", level=1)
            row = doc.add_table(rows=1, cols=2).rows[0]
            row.cells[0].text = "예금 잔액"
            row.cells[1].text = "[확인 필요]"
            doc.add_heading("검토·출처 대조표 (법원 서식과 분리)", level=1)
            appendix = doc.add_table(rows=1, cols=2).rows[0]
            appendix.cells[0].text = "예금 잔액"
            appendix.cells[1].text = "987654원"
            doc.save(path)
            body = quality.primary_body(path)
            self.assertTrue(body["appendix_excluded"])
            self.assertEqual(body["unresolved_placeholder_count"], 1)
            self.assertNotIn("987654", "\n".join(line["text"] for line in body["lines"]))
            artifacts, warnings = scalar.artifact_rows({"artifact_dir": directory, "artifacts": [{"name": "draft.docx"}]})
            self.assertEqual(warnings, [])
            self.assertEqual(scalar.artifact_field_status(artifacts, "bank_balance", "987654"), "value_not_found_with_label")

    def test_raw_extraction_sheet_does_not_count_as_populated_output(self):
        quality.PRIVATE.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=quality.PRIVATE, prefix="body_test_") as directory:
            path = Path(directory) / "raw.xlsx"
            workbook = Workbook()
            workbook.active.title = "추출값"
            workbook.active.append(["월소득", "987654"])
            workbook.save(path)
            artifacts, warnings = scalar.artifact_rows({"artifact_dir": directory, "artifacts": [{"name": "raw.xlsx"}]})
            self.assertEqual(warnings, [])
            self.assertEqual(scalar.artifact_field_status(artifacts, "monthly_income", "987654"), "field_label_not_located")

    def test_false_match_requires_exact_existing_evidence(self):
        check = {"scorable": True, "numeric_values": ["100000"]}
        lines = [{"index": 0, "text": "예금 잔액 | 99999원"}]
        verdict = {"status": "match", "line_indices": [0], "generated_quote": "예금 잔액 | 99999원", "reason_code": "supported"}
        self.assertEqual(quality.validate_verdict(check, verdict, lines)["status"], "different")
        verdict["generated_quote"] = "예금 잔액 | 100000원"
        self.assertEqual(quality.validate_verdict(check, verdict, lines)["status"], "missing")
        verdict["line_indices"] = [10]
        self.assertEqual(quality.validate_verdict(check, verdict, lines)["status"], "missing")

    def test_withheld_and_unscorable_are_never_matches(self):
        lines = [{"index": 0, "text": "예금 잔액 | [확인 필요]"}]
        verdict = {"status": "match", "line_indices": [0], "generated_quote": lines[0]["text"], "reason_code": "supported"}
        self.assertEqual(quality.validate_verdict({"scorable": True, "numeric_values": []}, verdict, lines)["status"], "missing")
        self.assertEqual(quality.validate_verdict({"scorable": False, "numeric_values": []}, verdict, lines)["status"], "target_unscorable")
        verdict["status"] = "target_unscorable"
        self.assertEqual(quality.validate_verdict({"scorable": True, "numeric_values": []}, verdict, lines)["status"], "missing")

    def test_prompt_values_are_not_recursively_expanded(self):
        result = quality.render_body("Value: {{x}}", {"x": "source {{secret}}", "secret": "should not appear"})
        self.assertEqual(result, "Value: source {{secret}}")


if __name__ == "__main__":
    unittest.main()
