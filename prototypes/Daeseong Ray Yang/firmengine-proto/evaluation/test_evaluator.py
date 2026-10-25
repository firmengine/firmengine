"""Synthetic regression checks for independent evaluation; no API calls."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
import evaluate_sample as evaluator


class EvaluationTests(unittest.TestCase):
    def test_exact_amounts_and_currency_units(self):
        self.assertTrue(evaluator.equal_values("1,234원", "1234"))
        self.assertFalse(evaluator.equal_values("1234", "12345"))
        self.assertTrue(evaluator.row_has_value("월소득 | 1,234원", "1234"))
        self.assertFalse(evaluator.row_has_value("월소득 | 12,345원", "1234"))

    def test_null_is_withheld_but_zero_is_a_value(self):
        self.assertIsNone(evaluator.find_field({"records": [{"key": "monthly_income", "value": None}]}, "monthly_income")[0])
        self.assertEqual(evaluator.find_field({"records": [{"key": "monthly_income", "value": "0"}]}, "monthly_income")[0], "0")

    def test_calculation_source_references_resolve_without_leaking_values(self):
        generated = {"records": [{"id": "record1", "key": "monthly_income", "sources": [{"document_id": "d1", "page": 2}]}],
                     "calculations": [{"key": "a", "sources": ["record1", "b"]}, {"key": "b", "sources": ["a", "monthly_income"]}]}
        self.assertEqual(evaluator.cited_pages(generated), {2})

    def test_duplicate_creditor_rows_do_not_double_count(self):
        row = {"key": "creditor_principal", "entity": "테스트은행", "value": "100"}
        self.assertEqual(evaluator.find_field({"records": [row, row]}, "debt_principal_total")[0], "100")
        self.assertIsNone(evaluator.find_field({"records": [row, dict(row, value="200")]}, "debt_principal_total")[0])
        self.assertIsNone(evaluator.find_field({"records": [row]}, "debt_total")[0])

    def test_reads_actual_office_artifacts_and_checks_label(self):
        evaluator.PRIVATE.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=evaluator.PRIVATE, prefix="synthetic_test_") as directory:
            path = Path(directory)
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "월소득"
            table.cell(0, 1).text = "1,234원"
            doc.save(path / "draft.docx")
            book = Workbook()
            book.active.append(["차량가액", 4321])
            book.save(path / "calc.xlsx")
            artifacts, warnings = evaluator.artifact_rows({"artifact_dir": directory, "artifacts": [{"name": "draft.docx"}, {"name": "calc.xlsx"}]})
            self.assertEqual(len(artifacts), 2)
            self.assertEqual(warnings, [])
            self.assertEqual(evaluator.artifact_field_status(artifacts, "monthly_income", "1234"), "value_present_with_label")
            self.assertEqual(evaluator.artifact_field_status(artifacts, "monthly_income", "12345"), "value_not_found_with_label")
            self.assertEqual(evaluator.artifact_field_status(artifacts, "vehicle_value", "4321"), "value_present_with_label")

    def test_reports_denominators_withheld_values_and_leakage(self):
        evaluator.PRIVATE.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=evaluator.PRIVATE, prefix="synthetic_test_") as directory:
            path = Path(directory)
            evaluator.write_json(path / "sample_manifest.json", {"source_sha256": "synthetic", "holdout_pages": [1], "evidence_pages": [2]})
            targets = []
            for key, value, basis in [("monthly_income", "1234", "direct"), ("vehicle_value", "1000", "direct"), ("liquidation_value", "6500", "direct"), ("client_name", None, "redacted")]:
                targets.append({"key": key, "value": value, "basis": basis, "page": 1, "source_label": key, "interpretation_sensitive": False})
            evaluator.write_json(path / "golden.json", {"source_sha256": "synthetic", "extraction_method": "synthetic", "field_targets": targets, "creditor_targets": [], "holdout_pages_used": [1]})
            evaluator.write_json(path / "generated.json", {"evidence_pages_used": [{"document_id": "d1", "page": 1}, {"document_id": "d1", "page": 2}], "records": [{"key": "monthly_income", "value": "1234", "sources": [{"page": 1}]}, {"key": "vehicle_value", "value": "999"}, {"key": "liquidation_value", "value": None}]})
            with patch.object(evaluator, "EVALUATION", path):
                report = evaluator.compare(path / "golden.json", path / "generated.json")
            self.assertEqual(report["scalar_comparison"]["counts"], {"match": 1, "different": 1, "withheld_or_missing": 1, "target_unscorable": 1})
            self.assertEqual(report["scalar_comparison"]["match_over_scorable_targets"], "1/3")
            self.assertEqual(report["scalar_comparison"]["match_over_present"], "1/2")
            self.assertEqual(report["leakage"]["heldout_page_citations_found"], [1])
            self.assertEqual(report["leakage"]["heldout_pages_in_input"], [1])
            self.assertEqual(report["artifacts"]["readable_count"], 0)
            serialized = json.dumps(report, ensure_ascii=False)
            self.assertNotIn('"expected_value"', serialized)
            self.assertNotIn('"generated_value"', serialized)


if __name__ == "__main__":
    unittest.main()
