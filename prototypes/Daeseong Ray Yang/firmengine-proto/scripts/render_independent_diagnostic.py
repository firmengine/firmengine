"""Offline original-renderer diagnostic; no approval or case mutation, no AI call."""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import shutil
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))

BANNER='진단 전용 · 정상 생성 보류 · 원본 채택 정책 유지 · 제출·검수 완료본 아님'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(run_name):
    from app.rendering import render_bundle
    from docx import Document
    from docx.shared import RGBColor
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    if not run_name.replace('-','').replace('_','').isalnum():
        raise ValueError('Invalid run name')
    directory=ROOT/'.runtime'/'evaluation'/'runs'/run_name
    source=directory/'generated.json'
    session=json.loads(source.read_text(encoding='utf-8'))
    before=hashlib.sha256(json.dumps(session,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    output=directory/'diagnostic'
    raw=output/'raw-renderer-output'
    marked=output/'artifacts'
    raw.mkdir(parents=True,exist_ok=True);marked.mkdir(parents=True,exist_ok=True)
    artifacts=render_bundle(deepcopy(session),raw)
    for artifact in artifacts:
        name=artifact['name']
        original=raw/name;target=marked/name
        shutil.copyfile(original,target)
        if artifact['format']=='docx':
            doc=Document(target)
            para=doc.paragraphs[0].insert_paragraph_before(BANNER) if doc.paragraphs else doc.add_paragraph(BANNER)
            for span in para.runs:
                span.bold=True;span.font.color.rgb=RGBColor.from_string('9C3027')
            doc.core_properties.comments=BANNER
            doc.save(target)
        elif artifact['format']=='xlsx':
            workbook=load_workbook(target)
            for sheet in workbook:
                sheet.insert_rows(1)
                sheet.cell(1,1,BANNER).font=Font(bold=True,color='9C3027')
            workbook.save(target)
        artifact.update(path=str(target),raw_renderer_sha256=sha(original),sha256=sha(target),diagnostic_only=True)
    after=hashlib.sha256(json.dumps(session,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    if before!=after:
        raise ValueError('Diagnostic changed the source snapshot')
    manifest=dict(run_id=run_name,created_at=datetime.now(timezone.utc).isoformat(),
                  source_snapshot_sha256=sha(source),source_snapshot_unchanged=True,
                  application_generation_success=False,case_database_changed=False,ai_calls=0,
                  original_assertion_acceptance_policy_preserved=True,
                  policy='Original renderer hides unverified/unaccepted assertions; this is not an all-candidate rendering.',
                  renderer_sha256=sha(ROOT/'app'/'rendering.py'),helper_sha256=sha(Path(__file__)),artifacts=artifacts)
    (output/'diagnostic-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(diagnostic_artifacts=len(artifacts),formats={kind:sum(a['format']==kind for a in artifacts) for kind in ('docx','xlsx')},
                         original_acceptance_policy_preserved=True,normal_generation_success=False),ensure_ascii=True))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--run-name',default='original-independent-20261026-001')
    run(parser.parse_args().run_name)
