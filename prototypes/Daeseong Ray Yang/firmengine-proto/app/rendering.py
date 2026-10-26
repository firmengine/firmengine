"""Editable court-form drafts with explicit provenance and review state.

This module consumes the application contract, never the held-out submitted
answers.  It renders version-pinned reconstructed DOCX templates whose sections
are mapped to the public court originals distributed in ``templates/originals``.
It does not make legal findings, infer missing amounts, sign, or select consent.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import MONEY_KEYS


TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates"
MISSING = "[확인 필요]"
CONFLICT = "[불일치·확인 필요]"
PERIOD = "[기간 선택 필요]"
TOKEN = re.compile(r"\{\{([a-zA-Z0-9_.-]+)\}\}")
XML_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
STATES = {
    "confirmed": "담당자 확인", "unreviewed": "미확인 추출값",
    "conflict": "불일치", "missing": "확인 필요", "calculated": "산술 계산",
    "pending": "계산 보류", "proposed": "검토 전 제안", "excluded": "검토 제외", "rejected": "제외", "superseded": "대체됨",
}
CALC_ALIASES = {
    "total_principal": "debt_principal_total",
    "total_interest": "debt_interest_total",
    "monthly_business_income": "monthly_business_income",
    "monthly_disposable_income": "monthly_disposable_income",
    "monthly_payment": "monthly_payment",
    "total_repayment": "total_repayment",
    "repayment_rate": "repayment_rate",
}
AMOUNT_KEYS = MONEY_KEYS
REPEAT_FIELDS = {
    'creditor': {'creditor_' + name: name for name in (
        'principal', 'interest', 'address', 'cause', 'phone', 'fax', 'number', 'contract_id',
        'content', 'loan_date', 'original_principal', 'principal_basis', 'interest_basis',
        'claim_type', 'priority_basis', 'due_date', 'collateral', 'secured_recovery',
        'shortfall', 'secured_amount', 'disputed_amount', 'undisputed_amount', 'dispute_reason',
        'litigation', 'assignment_order', 'other_details', 'claimed_principal', 'claimed_interest',
        'undisputed_interest', 'disputed_interest')},
    'asset': {'asset_' + name: name for name in (
        'value', 'description', 'kind', 'institution', 'account_number', 'seizure',
        'collateral', 'valuation_basis', 'ownership_share', 'return_amount', 'address', 'acquisition_date')},
    'month': {'monthly_revenue': 'revenue', 'monthly_expense': 'expense'},
    'family': {'family_' + name: name for name in (
        'relationship', 'age', 'cohabitation', 'occupation', 'monthly_income', 'assets', 'dependent')},
    'career': {'career_' + name: name for name in ('employer', 'business_type', 'position')},
    'enforcement': {'enforcement_' + name: name for name in ('kind', 'creditor', 'court', 'case_number')},
    'prior': {'prior_procedure_' + name: name for name in ('kind', 'court', 'status')},
}
REPEAT_ALIASES = {
    'secured': {'fields': {'collateral', 'secured_recovery', 'shortfall', 'secured_amount'},
                'types': {'secured', '별제권부', '담보부', '별제권부채권'}},
    'disputed': {'fields': {'dispute_reason', 'disputed_amount', 'litigation'},
                 'types': {'disputed', 'unconfirmed', '미확정', '다툼있는채권'}},
    'assignment': {'fields': {'assignment_order'}, 'types': set()},
    'other_claim': {'fields': {'other_details'}, 'types': set()},
    'priority': {'fields': {'priority_basis'}, 'types': {'priority', '우선권', '우선권있는채권'}},
    'estate': {'fields': set(), 'types': {'estate', '재단채권', '개인회생재단채권'}},
}
REPEAT_FIELDS.update({name: REPEAT_FIELDS['creditor'] for name in REPEAT_ALIASES})
REPEAT_FIELDS.update({name: {} for name in ('allocation', 'valuation', 'disposal', 'attachment', 'protected')})


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    return XML_CONTROLS.sub("", str(value))


def _amount(value: Any) -> str:
    """Format only unambiguous plain numeric strings; never coerce unknown to 0."""
    text = clean(value).strip()
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        return text
    try:
        number = Decimal(text)
        if number.is_finite():
            return format(number, ",f")
    except InvalidOperation:
        pass
    return text


def _records(session: dict, key: str | None = None) -> list[dict]:
    return [row for row in session.get("records", [])
            if (key is None or row.get("key") == key)
            and row.get("state") not in ("excluded", "rejected", "superseded")]


def _resolve_rows(rows: list[dict]) -> str:
    if not rows:
        return MISSING
    if any(row.get("state") == "conflict" for row in rows):
        return CONFLICT
    usable = [row for row in rows if row.get("value") is not None
              and clean(row.get("value")).strip() and row.get("state") != "missing"]
    if not usable:
        return MISSING
    confirmed = [row for row in usable if row.get("state") == "confirmed"]
    selected = confirmed or usable
    periods = {clean(row.get("period")).strip() for row in selected}
    if len(periods) > 1:
        return PERIOD
    values = {clean(row["value"]).strip() for row in selected}
    if len(values) != 1:
        return CONFLICT
    value = clean(selected[0]["value"]).strip()
    # A string of digits may be a phone/account/registration identifier. Only
    # semantic monetary fields receive thousands separators or Decimal parsing.
    return _amount(value) if all(row.get("key") in AMOUNT_KEYS for row in selected) else value


def _calculation(session: dict, key: str) -> str:
    rows = [row for row in session.get("calculations", []) if row.get("key") == key]
    if not rows or any(row.get("status") != "calculated" or row.get("value") is None for row in rows):
        proposed = [row for row in session.get('draft_calculations', [])
                    if row.get('key') == key and row.get('status') == 'proposed' and row.get('value') is not None]
        if proposed:
            rows = proposed
        else:
            return "[계산 보류]" if rows else MISSING
    values = {clean(row["value"]) for row in rows}
    return _amount(rows[0]["value"]) if len(values) == 1 else CONFLICT


def _review(session: dict, role: str) -> dict:
    value = session.get("reviews", {}).get(role)
    return value if isinstance(value, dict) else {}


def _review_valid(session: dict, role: str) -> bool:
    row = _review(session, role)
    return bool((row.get("name") or row.get("reviewer") or row.get("actor"))
                and row.get("revision") == session.get("revision"))


def review_status(session: dict) -> str:
    if (session.get("template_approved") is True
            and _review_valid(session, "caseworker") and _review_valid(session, "lawyer")):
        return "검토 완료본 · 실제 법원 제출 전"
    return "미검토 초안 · 담당자 및 변호사 확인 필요"


def _value(session: dict, key: str) -> str:
    if key == "review_status":
        return review_status(session)
    if key == "revision":
        return clean(session.get("revision", 0))
    if key == "session_id":
        return clean(session.get("id"))
    if key in CALC_ALIASES:
        return _calculation(session, CALC_ALIASES[key])
    records = _records(session, key)
    confirmed = [row for row in records if row.get('state') == 'confirmed']
    if confirmed:
        return _resolve_rows(confirmed)
    candidate = (session.get('draft_fields') or {}).get(key)
    if _supported_proposal(candidate) and candidate.get('value') is not None:
        return _resolve_rows([{'key': key, 'value': candidate['value'], 'state': 'unreviewed'}])
    return _resolve_rows(records)


def _supported_proposal(candidate) -> bool:
    return (isinstance(candidate, dict) and candidate.get('status') == 'proposed'
            and bool(candidate.get('source_record_ids') or candidate.get('source_pages')))


def _grouped(session: dict, prefix: str) -> list[dict]:
    if prefix in ('allocation', 'valuation', 'disposal', 'attachment', 'protected'):
        return _package_rows(session, prefix)
    if prefix in REPEAT_ALIASES:
        selection = REPEAT_ALIASES[prefix]
        return [row for row in _grouped(session, 'creditor')
                if any(row.get(key, MISSING) not in (MISSING, '', 'False', 'false', 'NONE', 'None', '없음', '해당없음') for key in selection['fields'])
                or re.sub(r'\s+', '', row.get('claim_type', '')).casefold() in selection['types']]
    keys = REPEAT_FIELDS[prefix]
    if prefix == 'creditor' and session.get('draft_creditors'):
        confirmed = [row for row in _records(session) if row.get('key') in ('creditor_principal', 'creditor_interest')]
        if not confirmed or not all(row.get('state') == 'confirmed' for row in confirmed):
            output = []
            for index, candidate in enumerate(session['draft_creditors'], 1):
                if not _supported_proposal(candidate):
                    continue
                row = {'entity': clean(candidate.get('name') or candidate.get('entity')) or MISSING,
                       'period': clean(candidate.get('period')) or MISSING}
                row.update({name: _resolve_rows([{'key': key, 'value': candidate.get(name), 'state': 'unreviewed'}])
                            for key, name in keys.items()})
                row['number'] = clean(candidate.get('number')) or str(index)
                row.update(monthly_payment=MISSING, total_payment=MISSING)
                output.append(row)
            return output
    grouped: dict[tuple[str, str], dict[str, list]] = {}
    display_entities = {}
    for record in _records(session):
        if record.get("key") not in keys:
            continue
        entity = clean(record.get("entity")).strip()
        # Whitespace alone does not create a different party/contract. Account
        # and contract identifiers remain part of the identity and are retained.
        identity = (clean(record.get('object_id') or record.get('contract_id')) or re.sub(r"\s+", "", entity).casefold(), clean(record.get("period")).strip())
        display_entities.setdefault(identity, entity)
        group = grouped.setdefault(identity, {})
        group.setdefault(keys[record["key"]], []).append(record)
    metadata_fields = {"creditor": set(keys.values()) - {'principal', 'interest', 'original_principal',
                           'secured_recovery', 'shortfall', 'secured_amount', 'disputed_amount', 'undisputed_amount'},
                       "asset": set(keys.values()) - {'value', 'return_amount'},
                       "month": set()}.get(prefix, set())
    for identity, fields in list(grouped.items()):
        entity, period = identity
        if not entity or period or not metadata_fields.intersection(fields):
            continue
        # An undated address/description describes an identified party or item;
        # it is not a new debt/asset snapshot. Never move undated monetary values.
        targets = [key for key, values in grouped.items()
                   if key[0] == entity and key[1] and set(values) - metadata_fields]
        if not targets:
            continue
        for target in targets:
            for name in metadata_fields.intersection(fields):
                grouped[target].setdefault(name, fields[name])
        if not (set(fields) - metadata_fields):
            del grouped[identity]
    output = []
    for index, ((entity, period), fields) in enumerate(grouped.items(), 1):
        row = {"entity": display_entities[(entity, period)] or MISSING, "period": period or MISSING}
        row.update({name: _resolve_rows(fields.get(name, [])) for name in keys.values()})
        if prefix == 'creditor' and row.get('number') == MISSING:
            row['number'] = str(index)
        # Allocation is a separate reviewed calculation, not proportional guessing.
        row.update(monthly_payment=MISSING, total_payment=MISSING)
        output.append(row)
    return output


def _decision(snapshot, key):
    row = (snapshot.get('package_decisions') or {}).get(key) or {}
    if (row.get('state') == 'VERIFIED' and row.get('revision') == snapshot.get('revision')
            and str(row.get('reviewer') or '').strip() and isinstance(row.get('value'), bool)):
        return row['value']
    return None


def _assertions(snapshot):
    rows = snapshot.get('assertions') or []
    replaced = {row.get('supersedes') for row in rows if row.get('supersedes')}
    return [row for row in rows if row.get('state') != 'RETRACTED' and row.get('id') not in replaced]


def _assertion_value(snapshot, key, object_id=None):
    rows = [row for row in _assertions(snapshot) if row.get('feature_key') == key
            and (object_id is None or row.get('object_id') == object_id)
            and (row.get('state') == 'DISPUTED' or row.get('accepted_for_filing') is True and row.get('state') == 'VERIFIED')]
    mapped = [{'key': '', 'value': row.get('value'), 'state': {'VERIFIED': 'confirmed', 'DISPUTED': 'conflict'}.get(row.get('state'), 'unreviewed'),
               'period': row.get('valid_at') or row.get('source_period') or ''} for row in rows]
    return _resolve_rows(mapped)


def _render_session(snapshot):
    """Canonical current assertions override their legacy fields without losing the audit copy."""
    if not snapshot.get('assertions'):
        return snapshot
    from .case_engine import FLAT_FEATURES
    result = deepcopy(snapshot)
    result['legacy_records'] = deepcopy(snapshot.get('records') or [])
    reverse = {}
    for legacy, feature in FLAT_FEATURES.items():
        reverse.setdefault(feature, []).append(legacy)
    reverse.update({'debts[].security_status': ['creditor_claim_type'],
                    'debts[].assignment_order': ['creditor_assignment_order'],
                    'debts[].disputed_amount': ['creditor_disputed_amount'],
                    'debts[].undisputed_amount': ['creditor_undisputed_amount'],
                    'debts[].collateral_description': ['creditor_collateral'],
                    'debts[].secured_recovery': ['creditor_secured_recovery'],
                    'debts[].shortfall': ['creditor_shortfall']})
    objects = {obj['id']: obj for obj in snapshot.get('objects', [])}
    mapped = []
    for row in _assertions(snapshot):
        feature = row.get('feature_key', '')
        keys = reverse.get(feature, [feature[5:]] if feature.startswith('form.') else [])
        for key in keys:
            if key.startswith('annual_') and row.get('measurement_period') != 'ANNUAL':
                continue
            if key in ('monthly_revenue', 'monthly_expense') and row.get('measurement_period') != 'MONTHLY':
                continue
            oid = row.get('object_id')
            obj = objects.get(oid, {})
            adopted = row.get('state') == 'VERIFIED' and row.get('accepted_for_filing') is True
            mapped.append({'id': row.get('id'), 'key': key, 'value': row.get('value') if adopted else None,
                'entity': obj.get('label') or row.get('object_label') or oid or '',
                'object_id': oid, 'unit': row.get('unit'),
                'period': row.get('source_period') or row.get('valid_at') or '',
                'state': 'confirmed' if adopted else 'conflict' if row.get('state') == 'DISPUTED' else 'missing',
                'sources': [{'document_id': ref.get('document_id'), 'page': ref.get('page'), 'quote': ref.get('excerpt')}
                            for ref in row.get('evidence_refs', [])]})
    mapped_keys = {row['key'] for row in mapped}
    result['records'] = [row for row in result.get('records', []) if row.get('key') not in mapped_keys] + mapped
    calculations = snapshot.get('calculation_results') or {}
    if isinstance(calculations, dict) and calculations.get('case_revision') == snapshot.get('revision'):
        aliases = {'monthly_net': 'monthly_business_income', 'total_nominal_repayment': 'total_repayment'}
        mapped_calcs = [{'key': aliases.get(key, key), 'value': value, 'status': 'calculated', 'unit': 'KRW',
                         'formula': row.get('rule_id'), 'sources': row.get('input_assertion_ids', [])}
                        for row in calculations.get('results', []) if row.get('outcome') in ('CALCULATED', 'PASS')
                        for key, value in row.get('values', {}).items()]
        names = {row['key'] for row in mapped_calcs}
        result['calculations'] = [row for row in result.get('calculations', []) if row.get('key') not in names] + mapped_calcs
    return result


def _package_rows(session, prefix):
    if prefix == 'allocation':
        current_ids = {row.get('id') for row in _assertions(session)}
        rows = [row for row in session.get('plan_allocations', [])
                if row.get('case_revision', row.get('revision')) == session.get('revision')
                and row.get('status') in ('calculated', 'verified', 'proposed')
                and row.get('source_assertion_ids') and set(row['source_assertion_ids']) <= current_ids]
        if not rows:
            return [dict(row, creditor_name=row['entity'], status='배분 근거 확인 필요') for row in _grouped(session, 'creditor')]
        money_fields = {'claim_confirmed', 'claim_unconfirmed', 'monthly_confirmed', 'monthly_reserved',
                        'total_confirmed', 'total_reserved', 'disposal_confirmed', 'disposal_reserved'}
        return [{**{key: _amount(value) if key in money_fields and value is not None else clean(value) if value is not None else MISSING
                    for key, value in row.items()}, 'source_ids': clean(row['source_assertion_ids'])} for row in rows]
    if prefix == 'valuation':
        keys = {'owner':'assets[].legal_owner', 'share':'assets[].ownership_share', 'as_of':'assets[].valuation_as_of',
                'method':'assets[].valuation_method', 'gross':'assets[].gross_value', 'secured_balance':'assets[].secured_balance',
                'protection_basis':'assets[].protection_basis', 'liquidation_contribution':'assets[].liquidation_contribution'}
        from .case_engine import ASSET_TYPES
        rows = []
        for obj in session.get('objects', []):
            if obj.get('type') not in ASSET_TYPES:
                continue
            row = {'label': clean(obj.get('label')) or obj['id'], 'evidence': clean([ref for item in _assertions(session)
                   if item.get('object_id') == obj['id'] for ref in item.get('evidence_refs', [])]) or MISSING}
            row.update({key: _assertion_value(session, feature, obj['id']) for key, feature in keys.items()})
            rows.append(row)
        return rows or [dict(row, label=row['entity'], gross=row.get('value'), method=row.get('valuation_basis'),
                            share=row.get('ownership_share'), as_of=row['period']) for row in _grouped(session, 'asset')]
    if prefix == 'disposal':
        rows = [row for row in session.get('disposals', []) if row.get('state') != 'RETRACTED']
        for item in _assertions(session):
            if item.get('feature_key') == 'assets[].disposal_history[]':
                value = item.get('value')
                additions = value if isinstance(value, list) else [value] if isinstance(value, dict) else []
                rows.extend({**row, 'source_ids': [item.get('id')]} for row in additions if isinstance(row, dict))
        return [{key: clean(value) if value is not None else MISSING for key, value in row.items()} for row in rows]
    if prefix == 'protected':
        return [{key: clean(value) if value is not None else MISSING for key, value in row.items()}
                for row in session.get('protected_assets', []) if row.get('state') == 'VERIFIED']
    documents = {doc['id']: doc for doc in session.get('documents', [])}
    selected = {(row.get('requirement_id'), row.get('document_id')): row for row in session.get('court_attachments', [])
                if row.get('included') is True and row.get('revision') == session.get('revision')}
    output = []
    for requirement in _case_requirements(session):
        for doc_id in requirement.get('matched_document_ids') or [None]:
            doc = documents.get(doc_id, {})
            output.append({'number': str(len(output)+1), 'title': doc.get('name') or requirement.get('label') or MISSING,
                'purpose': requirement.get('label') or requirement.get('id'), 'subject': requirement.get('subject_id') or '신청인',
                'period': clean(requirement.get('period')), 'included': '포함 확인' if (requirement.get('requirement_id'), doc_id) in selected else '편철 확인 필요',
                'pages': clean(doc.get('pages')) or MISSING, 'document_id': doc_id or MISSING})
    return output


def _case_requirements(snapshot):
    return snapshot.get('requirements') or (snapshot.get('evaluation') or {}).get('requirements') or (snapshot.get('assessment_result') or {}).get('requirements') or []


def _package_applicability(snapshot, entry):
    rule = entry['required_when']
    if rule == 'always':
        return True
    if rule == 'income_only_plan':
        return _decision(snapshot, 'plan_asset_disposal') is not True
    if rule.startswith('decision:'):
        return _decision(snapshot, rule.split(':', 1)[1])
    if rule.startswith('creditor:'):
        return bool(_grouped(snapshot, rule.split(':', 1)[1]))
    if rule == 'assets_present':
        return bool(_package_rows(snapshot, 'valuation'))
    return None


def court_package(snapshot: dict) -> list[dict]:
    """Evaluate court OUTPUT forms, attached evidence and signatures separately.

    A generated file is not automatically content-complete or signed. Missing
    output files block filing completeness, not generation itself. Unresolved
    optional applications are never silently selected as universal attachments.
    """
    snapshot = _render_session(snapshot)
    definitions = json.loads((TEMPLATE_DIR.parent/'rules/court_package.json').read_text(encoding='utf-8'))
    forms = {row['code']: row for row in _template_catalog()['forms']}
    artifacts = {row.get('template_id'): row for row in snapshot.get('artifacts', [])
                 if row.get('revision') == snapshot.get('revision')}
    checks = []
    for entry in definitions['entries']:
        applicable = _package_applicability(snapshot, entry)
        required = applicable is True and entry.get('mandatory_when_applicable', True)
        missing = []
        for key in entry['required_features']:
            if key == 'plan_allocations':
                rows = _package_rows(snapshot, 'allocation')
                if not rows or any(row.get('status') not in ('calculated', 'verified') or not all(row.get(name) not in (None, '', MISSING) for name in
                    ('claim_confirmed','claim_unconfirmed','monthly_confirmed','monthly_reserved','total_confirmed','total_reserved')) for row in rows):
                    missing.append(key)
            elif key == 'asset_valuations':
                rows = _package_rows(snapshot, 'valuation')
                if not rows or any(any(row.get(key) in (None, '', MISSING, CONFLICT) for key in
                    ('owner','share','as_of','method','gross','secured_balance','protection_basis','liquidation_contribution')) for row in rows):
                    missing.append(key)
            elif key == 'protected_assets':
                if not _package_rows(snapshot, 'protected'):
                    missing.append(key)
            elif key.startswith('creditor_'):
                prefix = entry['required_when'].split(':', 1)[1] if entry['required_when'].startswith('creditor:') else 'creditor'
                rows = _grouped(snapshot, prefix)
                field = key[len('creditor_'):]
                if not rows or any(row.get(field) in (None, '', MISSING, CONFLICT, PERIOD) for row in rows):
                    missing.append(key)
            elif _value(snapshot, key) in (MISSING, CONFLICT, PERIOD, '[계산 보류]'):
                missing.append(key)
        form = forms.get(entry['template_id'])
        present = entry['template_id'] in artifacts
        review = (snapshot.get('package_reviews') or {}).get(entry['id']) or {}
        reviewed = bool(review.get('revision') == snapshot.get('revision') and review.get('state') == 'VERIFIED'
                        and str(review.get('reviewer') or '').strip())
        status = ('NOT_APPLICABLE' if applicable is False else 'DECISION_REQUIRED' if applicable is None
                  else 'TEMPLATE_UNAVAILABLE' if not form else 'NOT_GENERATED' if not present
                  else 'CONTENT_INCOMPLETE' if missing else 'REVIEW_REQUIRED' if not reviewed else 'CONTENT_REVIEWED')
        checks.append({**deepcopy(entry), 'applicable': applicable if applicable is not None else 'unknown',
            'required': required, 'template_available': bool(form), 'missing_features': missing,
            'artifact_present': present, 'artifact_name': artifacts.get(entry['template_id'], {}).get('name'),
            'status': status, 'case_revision': snapshot.get('revision'),
            'blocks_filing': required and status != 'CONTENT_REVIEWED',
            'pre_generation_blocker': required and (not form or bool(missing)),
            'content_scope': 'Minimum supported fields plus explicit document review; not a full court completeness certification.'})
        if required and form and form['code'].startswith('D'):
            signature = (snapshot.get('package_signatures') or {}).get(entry['id']) or {}
            signed = bool(signature.get('revision') == snapshot.get('revision') and signature.get('state') == 'VERIFIED'
                          and str(signature.get('reviewer') or '').strip()
                          and signature.get('signed_document_id') in {doc['id'] for doc in snapshot.get('documents', [])})
            checks.append({'id':'signature:'+entry['id'],'title':entry['title']+' 서명·날인 확인',
                'kind':'MANUAL_SIGNATURE','required_when':entry['required_when'],'applicable':True,'required':True,
                'depends_on':[entry['id']],'source_ids':entry['source_ids'],'template_id':None,
                'status':'VERIFIED' if signed else 'SIGNATURE_REQUIRED','blocks_filing':not signed,
                'pre_generation_blocker':False,'case_revision':snapshot.get('revision')})
    documents = {doc['id']:doc for doc in snapshot.get('documents', [])}
    for req in _case_requirements(snapshot):
        ids = req.get('matched_document_ids', [])
        included = [row.get('document_id') for row in snapshot.get('court_attachments', [])
                    if row.get('requirement_id') == req.get('requirement_id') and row.get('included') is True
                    and row.get('revision') == snapshot.get('revision') and row.get('document_id') in ids
                    and row.get('document_id') in documents]
        alternatives = req.get('alternatives') or []
        if alternatives:
            from .case_engine import _alternatives_met
            included_types = {documents[doc_id].get('document_type') for doc_id in included}
            unavailable = bool(snapshot.get('primary_income_document_unavailable'))
            object_metadata = next((obj.get('metadata') or {} for obj in snapshot.get('objects', []) if obj['id'] == req.get('subject_id')), {})
            unavailable = unavailable or bool(object_metadata.get('primary_unavoidably_unavailable'))
            complete_group = _alternatives_met(alternatives, included_types, unavailable)
        else:
            complete_group = bool(included)
        ready = req.get('status') == 'SATISFIED' and complete_group
        checks.append({'id':'evidence:'+req['requirement_id'],'title':req.get('label') or req.get('id'),
            'kind':'EVIDENCE_ATTACHMENT','required_when':'applicable_requirement','applicable':True,'required':True,
            'source_ids':req.get('legal_basis', []),'depends_on':[],'template_id':None,
            'document_ids':included,'available_document_ids':ids,'required_features':req.get('purpose_features', []),
            'status':'INCLUDED' if ready else 'ATTACHMENT_NOT_INCLUDED' if req.get('status')=='SATISFIED' else 'EVIDENCE_UNRESOLVED',
            'blocks_filing':not ready,'pre_generation_blocker':False,'case_revision':snapshot.get('revision')})
    for attachment in definitions.get('conditional_attachments', []):
        if _decision(snapshot, attachment['decision']) is not True:
            continue
        identifier = 'package:' + attachment['id']
        available = {doc_id for doc_id, doc in documents.items() if doc.get('document_type') in attachment['document_types']}
        included = [row['document_id'] for row in snapshot.get('court_attachments', [])
                    if row.get('requirement_id') == identifier and row.get('document_id') in available
                    and row.get('included') is True and row.get('revision') == snapshot.get('revision')]
        review = (snapshot.get('package_reviews') or {}).get(identifier) or {}
        reviewed = bool(review.get('state') == 'VERIFIED' and review.get('revision') == snapshot.get('revision')
                        and str(review.get('reviewer') or '').strip())
        ready = bool(included) and reviewed
        checks.append({'id':identifier,'title':attachment['title'],'kind':'EVIDENCE_ATTACHMENT',
            'required_when':'decision:'+attachment['decision'],'applicable':True,'required':True,
            'source_ids':attachment['source_ids'],'depends_on':[attachment['decision']],'template_id':None,
            'document_ids':included,'available_document_ids':sorted(available),
            'status':'INCLUDED' if ready else 'REVIEW_REQUIRED' if included else 'ATTACHMENT_NOT_INCLUDED',
            'blocks_filing':not ready,'pre_generation_blocker':False,'case_revision':snapshot.get('revision')})
    return checks


def _paragraphs(document):
    def table_paragraphs(table):
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                yield from cell.paragraphs
                for nested in cell.tables:
                    yield from table_paragraphs(nested)
    yield from document.paragraphs
    for table in document.tables:
        yield from table_paragraphs(table)
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _replace_paragraph(paragraph, resolver):
    previous = paragraph.text
    replaced = TOKEN.sub(lambda match: clean(resolver(match.group(1))), previous)
    if replaced == previous:
        return
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replaced)


def _expand_tables(document, session):
    for table in document.tables:
        for row in list(table.rows):
            text = " ".join(cell.text for cell in row.cells)
            prefix = next((p for p in REPEAT_FIELDS if "{{" + p + "." in text), None)
            if not prefix:
                continue
            groups = _grouped(session, prefix) or [{"entity": MISSING, "period": MISSING}]
            index = table._tbl.index(row._tr)
            for offset, group in enumerate(groups):
                new_tr = deepcopy(row._tr)
                table._tbl.insert(index + offset, new_tr)
                for paragraph in new_tr.iter(qn("w:p")):
                    # Paragraph.text preserves Word line-break elements between
                    # fields; joining only w:t nodes would concatenate values.
                    _replace_paragraph(Paragraph(paragraph, table),
                        lambda key: group.get(key.split('.', 1)[1], MISSING)
                        if key.startswith(prefix + '.') else '{{' + key + '}}')
            table._tbl.remove(row._tr)


def _template_catalog() -> dict:
    catalog = json.loads((TEMPLATE_DIR / "catalog.json").read_text(encoding="utf-8"))
    for entry in catalog["forms"]:
        for asset in [entry["template"], *entry["originals"]]:
            path = TEMPLATE_DIR / asset["path"]
            if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != asset["sha256"]:
                raise ValueError(f"Template asset changed; update mapping and approval: {asset['path']}")
    return catalog


def _set_cell(cell, value):
    cell.text = clean(value)


def _add_appendix(document, session: dict, form: dict, generated_at: str):
    document.add_page_break()
    document.add_heading("검토·출처 대조표 (법원 서식과 분리)", level=1)
    metadata = [
        ("상태", review_status(session)), ("생성시각", generated_at),
        ("사건 세션 / 입력 버전", f"{session.get('id', '')} / {session.get('revision', 0)}"),
        ("서식", f"{form['code']} / {form['template_version']}"),
        ("원본 대응", "공식 원본의 항목·구역을 편집형으로 재구성. 동일한 페이지 배치가 아니며 양식 검토가 필요합니다."),
        ("서식 확인", "이 세션에서 확인됨" if session.get("template_approved") else "검토 대기"),
    ]
    for role, label in [("caseworker", "담당자"), ("lawyer", "변호사")]:
        row = _review(session, role)
        actor = row.get("name") or row.get("reviewer") or row.get("actor") or "확인 전"
        metadata.append((label, f"{actor} / {row.get('at', '')} / 확인 버전 {row.get('revision', '없음')}"))
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in metadata:
        cells = table.add_row().cells
        _set_cell(cells[0], label)
        _set_cell(cells[1], value)
    document.add_paragraph("서명·날인·동의 선택은 자동 작성하지 않습니다. 외부에서 편집한 파일은 앱의 확인 기록과 자동 동기화되지 않습니다.")
    document.add_heading("원본 양식 출처", level=2)
    for source in form["originals"]:
        document.add_paragraph(f"{source['title']}\n{source['url']}\nSHA-256: {source['sha256']}\n수집일: {source['retrieved_at']} (서식 시행일과 구분)")
    keys = set(form["record_keys"])
    relevant = [row for row in _records(session) if row.get("key") in keys]
    document.add_heading("작성 항목과 원문 근거", level=2)
    if not relevant:
        document.add_paragraph("연결된 증빙 항목이 없습니다. 빈칸을 임의의 0이나 사실로 채우지 않았습니다.")
    for row in relevant:
        label = row.get("label") or row.get("key")
        document.add_paragraph(f"{label} / {row.get('entity', '')} / {row.get('period', '')}", style="Heading 3")
        document.add_paragraph(f"값: {_resolve_rows([row])} · 상태: {STATES.get(row.get('state'), row.get('state', ''))} · 항목 ID: {row.get('id', '')}")
        for source in row.get("sources", []):
            excerpt = clean(source.get("quote"))
            if len(excerpt) > 700:
                excerpt = excerpt[:700] + " … [긴 인용은 앱 원문에서 확인]"
            document.add_paragraph(f"{source.get('document_id', '')} / {source.get('page', '쪽 미상')}쪽\n{excerpt}")
    document.add_heading("계산근거", level=2)
    for row in session.get("calculations", []):
        if row.get("key") not in form.get("calculation_keys", []):
            continue
        document.add_paragraph(f"{row.get('label') or row.get('key')}: {_calculation(session, row['key'])} {row.get('unit', '')}\n산식: {clean(row.get('formula'))}\n상태: {STATES.get(row.get('status'), row.get('status', ''))}")


def _sheet(workbook, title: str, headers: list[str], rows: list[list]):
    sheet = workbook.create_sheet(title)
    for row in [headers, *rows]:
        index = sheet.max_row + 1 if sheet.max_row > 1 or sheet.cell(1, 1).value is not None else 1
        for column, value in enumerate(row, 1):
            # Explicit text prevents formula injection and preserves monetary precision.
            cell = sheet.cell(index, column)
            cell.value = clean(value)[:32767]
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="183E35")
        cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 25 if index < len(headers) else 55
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=10)
    return sheet


def _workbook(session: dict, catalog: dict, generated_at: str) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "요약", ["항목", "내용"], [
        ["상태", review_status(session)], ["세션", session.get("id")],
        ["입력 버전", session.get("revision")], ["생성시각", generated_at],
        ["금액 형식", "정밀도 보존을 위해 원 단위 값을 텍스트로 저장. 산식은 실행 수식이 아닌 계산엔진의 근거 표현입니다."],
        ["편집 안내", "외부 편집 결과는 앱 확인 기록과 자동 동기화되지 않습니다."],
        ["기간", "동일 항목의 여러 기간은 개별 행으로 유지하며 합산하지 않습니다."],
    ])
    _sheet(workbook, "계산근거", ["키", "항목", "값", "단위", "상태", "산식", "출처"], [
        [row.get("key"), row.get("label"), row.get("value") if row.get("status") == "calculated" else MISSING,
         row.get("unit"), STATES.get(row.get("status"), row.get("status")), row.get("formula"), row.get("sources")]
        for row in session.get("calculations", [])
    ])
    _sheet(workbook, "추출값", ["ID", "키", "대상", "항목", "값", "단위", "기간", "상태", "신뢰도"], [
        [row.get("id"), row.get("key"), row.get("entity"), row.get("label"),
         row.get("value") if row.get("value") is not None else MISSING, row.get("unit"), row.get("period"),
         STATES.get(row.get("state"), row.get("state")), row.get("confidence")]
        for row in session.get("records", [])
    ])
    _sheet(workbook, "출처", ["항목 ID", "키", "문서", "페이지", "원문 발췌"], [
        [row.get("id"), row.get("key"), source.get("document_id"), source.get("page"), source.get("quote")]
        for row in session.get("records", []) for source in row.get("sources", [])
    ])
    _sheet(workbook, "채권자", ["채권자", "기준기간", "원금", "이자", "주소", "발생원인"], [
        [row.get(k, MISSING) for k in ["entity", "period", "principal", "interest", "address", "cause"]]
        for row in _grouped(session, "creditor")
    ])
    _sheet(workbook, "월별매출경비", ["대상 월", "자료 기간", "매출", "사업경비", "검토"], [
        [row["entity"], row["period"], row.get("revenue", MISSING), row.get("expense", MISSING), "인정값·대상기간은 담당자 확인 필요"]
        for row in _grouped(session, "month")
    ])
    _sheet(workbook, "재산", ["재산", "기준기간", "가액", "명세"], [
        [row.get(k, MISSING) for k in ["entity", "period", "value", "description"]]
        for row in _grouped(session, "asset")
    ])
    _sheet(workbook, "검토", ["역할", "확인자", "시각", "확인 버전", "현재 버전과 일치"], [
        [label, _review(session, role).get("name") or _review(session, role).get("reviewer"),
         _review(session, role).get("at"), _review(session, role).get("revision"),
         "예" if _review_valid(session, role) else "아니오"]
        for role, label in [("caseworker", "담당자"), ("lawyer", "변호사")]
    ])
    _sheet(workbook, "서식정보", ["코드", "서식명", "템플릿 버전", "상태", "원본 파일", "SHA-256", "출처 URL"], [
        [form["code"], form["title"], form["template_version"], "편집형 재구성·전문가 확인 필요",
         source["path"], source["sha256"], source["url"]]
        for form in catalog["forms"] for source in form["originals"]
    ])
    _sheet(workbook, '변제예정액표', ['채권번호','채권자','확정 원금','미확정 원금','월 변제액','월 유보액','총 변제액','총 유보액','재산처분 배분','재산처분 유보','상태·출처'], [
        [row.get(key, MISSING) for key in ('number','creditor_name','claim_confirmed','claim_unconfirmed','monthly_confirmed','monthly_reserved','total_confirmed','total_reserved','disposal_confirmed','disposal_reserved')]
        + [clean([row.get('status'), row.get('source_ids')])] for row in _package_rows(session, 'allocation')])
    _sheet(workbook, '법원제출패키지', ['ID','종류','서류','해당성','필수','상태','누락 항목','제출 보류'], [
        [row.get(key) for key in ('id','kind','title','applicable','required','status','missing_features','blocks_filing')]
        for row in court_package(session)])
    _sheet(workbook, '편철증빙', ['자료명','목적','대상','기간','포함 여부','문서 ID'], [
        [row.get(key, MISSING) for key in ('title','purpose','subject','period','included','document_id')]
        for row in _package_rows(session, 'attachment')])
    if session.get('assertions'):
        _sheet(workbook, '사건주장근거', ['ID','객체','Feature','값','상태','기준일','근거'], [
            [row.get(key) for key in ('id','object_id','feature_key','value','state','valid_at','evidence_refs')]
            for row in session['assertions']])
    _sheet(workbook, '보정위험체크', ['위험 ID','확인 항목','해당성','결과','생성 차단','필수 피처','근거','연결 검토'], [
        [row.get(key) for key in ('risk_id','title','applicable','outcome','blocks_generation','required_features','source_ids','finding_ids')]
        for row in session.get('risk_checklist', [])])
    _sheet(workbook, '위험이력', ['위험 ID','항목','이전 결과','현재 결과','사건 버전','기록 시각'], [
        [row.get(key) for key in ('risk_id','title','previous','outcome','case_revision','recorded_at')]
        for row in session.get('risk_history', [])])
    _sheet(workbook, '추가증빙요청', ['요청 ID','대상','목적','기간','상태','누락 피처','대체 증빙','법적 근거'], [
        [row.get(key) for key in ('requirement_id','subject_id','label','period','status','missing_features','alternatives','legal_basis')]
        for row in session.get('requirements', [])])
    _sheet(workbook, 'n8n계산', ['규칙','버전','결과','값','입력 주장 ID','누락 피처','비고'], [
        [row.get(key) for key in ('rule_id','rule_version','outcome','values','input_assertion_ids','missing_features','notes')]
        for row in (session.get('calculation_results') or {}).get('results', [])])
    return workbook


def render_bundle(session: dict, output_dir: Path) -> list[dict]:
    """Write the applicable editable court-output package and its evidence workbook."""
    session = _render_session(session)
    catalog = _template_catalog()
    definitions = json.loads((TEMPLATE_DIR.parent/'rules/court_package.json').read_text(encoding='utf-8'))
    selected = {row['template_id']: row for row in definitions['entries'] if _package_applicability(session, row) is True}
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    generated_at = datetime.now(timezone.utc).isoformat()
    revision = int(session.get("revision", 0))
    artifacts = []
    for form in catalog["forms"]:
        if form['code'] not in selected:
            continue
        index = len(artifacts) + 1
        document = Document(TEMPLATE_DIR / form["template"]["path"])
        for paragraph in _paragraphs(document):
            _replace_paragraph(paragraph, lambda key: "{{" + key + "}}"
                               if key.startswith(tuple(prefix + '.' for prefix in REPEAT_FIELDS)) else _value(session, key))
        # Expand user-provided repeat values last; literal {{...}} in evidence
        # must never be interpreted as another template instruction.
        _expand_tables(document, session)
        _add_appendix(document, session, form, generated_at)
        document.core_properties.title = form["title"]
        document.core_properties.subject = review_status(session)
        document.core_properties.author = "Firm Engine local document workflow"
        name = f"{index:02d}_{form['title'].replace(' ', '')}_r{revision}.docx"
        document.save(output_dir / name)
        artifacts.append({"name": name, "title": form["title"], "kind": form["kind"], "format": "docx",
                          'template_id':form['code'], 'package_id':selected[form['code']]['id'], 'revision':revision})
    name = f"{len(artifacts)+1:02d}_계산근거와검토대조표_r{revision}.xlsx"
    _workbook({**session, 'artifacts': artifacts}, catalog, generated_at).save(output_dir / name)
    artifacts.append({"name": name, "title": "계산근거와 검토 대조표", "kind": "workbook", "format": "xlsx",
                      'template_id':'WORKBOOK','package_id':'review_workbook','revision':revision})
    return artifacts
