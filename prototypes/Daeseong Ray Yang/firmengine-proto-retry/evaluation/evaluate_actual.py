"""Strict independent body/repeated-row coverage against a frozen private truth.

This evaluator never sends its answers to extraction/generation. It compares actual
DOCX bodies, not the generator's JSON or calculation appendix. It reports normal
generation blocks separately from diagnostic candidate rendering and quality.
"""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from docx import Document

VERSION = "strict-body-and-row-coverage-2-label-only-fix"


def normalize(value):
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = re.sub(r"(\d{4})\s*[년.\-/]\s*(\d{1,2})\s*[월.\-/]\s*(\d{1,2})\s*일?", lambda m:f"{int(m[1]):04d}{int(m[2]):02d}{int(m[3]):02d}", text)
    text = text.replace("주식회사", "").replace("(주)", "").replace("㈜", "")
    text = re.sub(r"(?<=\d),(?=\d)", "", text)
    text = re.sub(r"(?<=\d)\s*원", "", text)
    return re.sub(r"[^0-9A-Za-z가-힣.]", "", text).lower()


def numeric_values(text):
    text = re.sub(r"(?<=\d),(?=\d)", "", str(text))
    result=[]
    for token in re.findall(r"(?<![\d.])-?\d+(?:\.\d+)?(?![\d.])", text):
        try: result.append(Decimal(token))
        except InvalidOperation: pass
    return result


def matches(value, context):
    expected=normalize(value)
    if not expected:
        return False
    if re.fullmatch(r"[\d,.\s]+(?:원|%|개월|명|회)?", str(value).strip()):
        values=numeric_values(value)
        actual=numeric_values(context)
        return bool(values) and all(number in actual for number in values)
    return expected in normalize(context)


def label_only(check):
    """A printed heading with empty value cells is not a reproduced case fact."""
    if check["type"] == "field_coverage":
        return normalize(check.get("expected")) == normalize(check.get("label"))
    if check["type"] == "repeated_row_coverage":
        nonempty=[str(value).strip() for value in check["expected"] if str(value).strip()]
        expense_labels={"생계비","주거비","의료비","교육비","교통비","식비","합계","계"}
        return len(nonempty)==1 and normalize(nonempty[0]) in expense_labels
    return False


def aliases(check):
    label=check.get("label", "")
    base=normalize(label).replace("의", "").replace("평균", "").replace("조정된", "인정")
    values=[base]
    pairs=[
        (("신청인성명","채무자성명","신청인예금주"),["신청인","채무자","성명"]),
        (("월변제예정금액","월변제예정액"),["월변제예정액","월변제금"]),
        (("변제예정기간","변제기간"),["변제기간","변제횟수","변제기간개월"]),
        (("채권현재액총합계","무담보회생채권액합계"),["채권현재액합계","총채무"]),
        (("원금합계",),["원금합계","채무원금합계"]),
        (("이자합계",),["이자합계","채무이자합계"]),
        (("조정생계비","인정생계비","월생계비","월평균생계비"),["인정생계비","월생계비","생계비"]),
        (("월실제가용소득","월가용소득"),["월가용소득","가용소득"]),
        (("월수입","월평균수입"),["월순소득","월순수입","월평균순수입","월평균순소득"]),
        (("총실제가용소득","총변제예정액"),["명목총변제예정액","총변제예정액"]),
        (("청산가치",),["청산가치","청산가치합계"]),
        (("현재가치",),["현재가치"]),
        (("제1회변제일","변제예정시작일","변제기간시작일"),["최초납입예정일","변제기간"]),
        (("이후변제일",),["매월납입일","매월변제일","납입일"]),
    ]
    for triggers, alternatives in pairs:
        if any(normalize(trigger) in base for trigger in triggers):
            values.extend(normalize(item) for item in alternatives)
    if base == "채무자":
        values.extend(["신청인","채무자","성명"])
    return [value for value in values if value]


def read_docx(path):
    doc=Document(path)
    paragraphs=[paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    contexts=paragraphs.copy()
    rows=[]
    for ti,table in enumerate(doc.tables):
        headers=[cell.text for cell in table.rows[0].cells] if table.rows else []
        for ri,row in enumerate(table.rows):
            cells=[cell.text for cell in row.cells]
            text=" | ".join(cells)
            contexts.append(text)
            for index,cell in enumerate(cells):
                if index<len(headers):
                    contexts.append(headers[index]+": "+cell)
            rows.append({"id":f"table-{ti}-row-{ri}","cells":cells,"text":text})
    return {"contexts":contexts,"rows":rows,"sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"paragraphs":paragraphs}


def evaluate(output_dir: Path | None, truth_path: Path, *, mode="normal", workflow_status="unknown", expected_truth_sha256=None):
    raw=truth_path.read_bytes()
    digest=hashlib.sha256(raw).hexdigest()
    if expected_truth_sha256 and digest!=expected_truth_sha256:
        raise ValueError("고정 독립 기준답의 SHA-256이 달라졌습니다.")
    truth=json.loads(raw)
    if truth["metadata"].get("generated_outputs_read_before_freeze"):
        raise ValueError("생성물을 보기 전에 고정한 독립 기준답이 필요합니다.")
    directory=Path(output_dir).resolve() if output_dir else None
    documents={}
    if directory and directory.exists():
        for path in directory.glob("*.docx"):
            documents[path.name]=read_docx(path)
    results=[]
    used_rows={}
    for check in truth["checks"]:
        result={key:value for key,value in check.items() if key not in ("source_quote",)}
        if check.get("unscorable"):
            results.append(result | {"status":"unscorable"})
            continue
        if label_only(check):
            results.append(result | {"status":"unsupported","reason":"고정 기준의 양식 라벨만 있고 대조할 사건 기재값이 없어 일치로 채점하지 않음"})
            continue
        targets=[(name,documents[name]) for name in check.get("files",[]) if name in documents]
        if not targets:
            results.append(result | {"status":"missing_document","reason":"평가 대상 법원 서식 본문 파일이 없음"})
            continue
        if check["type"]=="field_coverage":
            bound=[]
            for name,doc in targets:
                for context in doc["contexts"]:
                    candidate=normalize(context).replace("의", "").replace("평균", "").replace("조정된", "인정")
                    if any(alias in candidate for alias in aliases(check)):
                        bound.append((name,context))
            hit=next(((name,context) for name,context in bound if matches(check["expected"],context)),None)
            if hit:
                result|={"status":"match","actual_file":hit[0],"actual_context":hit[1]}
            elif bound:
                result|={"status":"different","actual_contexts":[context for _,context in bound[:4]]}
            else:
                result|={"status":"missing","reason":"해당 본문 필드 또는 항목에 연결된 값 없음"}
        elif check["type"]=="repeated_row_coverage":
            expected=[str(value) for value in check["expected"] if str(value).strip() and not re.search(r"가림|masked|unreadable|판독\s*불가|\[.*가려",str(value),re.I)]
            if not expected:
                result|={"status":"unscorable","reason":"반복행에서 대조할 판독 가능한 셀이 없음"}
            else:
                key=(check["source_page"],check["table_title"])
                used=used_rows.setdefault(key,set())
                hit=None
                for name,doc in targets:
                    for row in doc["rows"]:
                        identity=(name,row["id"])
                        if identity not in used and all(matches(value,row["text"]) for value in expected):
                            hit=(name,row,identity)
                            break
                    if hit: break
                if hit:
                    used.add(hit[2])
                    result|={"status":"match","actual_file":hit[0],"actual_cells":hit[1]["cells"]}
                else:
                    result|={"status":"missing","reason":"동일 반복행 안에서 원본의 기재값·대상·계약별 항목이 함께 재현되지 않음"}
        else:
            result|={"status":"technical_error","reason":"지원되지 않는 기준답 유형"}
        results.append(result)
    counts=Counter(row["status"] for row in results)
    grouped={}
    for kind in sorted({row["document_type"] for row in results}):
        grouped[kind]=dict(Counter(row["status"] for row in results if row["document_type"]==kind))
    return {
        "evaluated_at":datetime.now(timezone.utc).isoformat(),"comparator_version":VERSION,"mode":mode,
        "truth_sha256":digest,"source_sha256":truth["metadata"]["source_sha256"],
        "workflow":{"status":workflow_status,"counted_as_generation_success":mode=="normal" and workflow_status=="completed" and bool(documents),
                    "diagnostic_only":mode=="diagnostic"},
        "coverage":{"criteria":len(results),"pre_unscorable":sum(bool(check.get("unscorable")) for check in truth["checks"]),
                    "scorable":len(results)-counts["unscorable"]-counts["unsupported"],"summary":dict(counts),"by_document_type":grouped,
                    "generated_document_count":len(documents),"document_hashes":{name:doc["sha256"] for name,doc in documents.items()}},
        "limitations":["독립 원본 전사 기준은 전문가 인증 정답이 아닙니다.","정규화 본문 필드 및 반복행 기재값 범위 비교이며 의미가 동등한 서술의 법적 적정성을 자동 인증하지 않습니다.","가림·판독불가 항목은 채점하지 않습니다.","판정기 v2는 값 없는 인쇄 라벨의 거짓 일치를 수정했으며 원래442개 기준과 기대값은 변경하지 않았습니다. 모든 실행에 같은 수정본을 적용합니다.","진단 후보 문서는 정상 생성 또는 생성 후 HITL 완료로 세지 않습니다."],
        "results":results}


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--output-dir",type=Path)
    parser.add_argument("--truth",type=Path,required=True)
    parser.add_argument("--truth-sha256")
    parser.add_argument("--report",type=Path,required=True)
    parser.add_argument("--mode",choices=("normal","diagnostic"),default="normal")
    parser.add_argument("--workflow-status",default="unknown")
    args=parser.parse_args()
    report=evaluate(args.output_dir,args.truth,mode=args.mode,workflow_status=args.workflow_status,expected_truth_sha256=args.truth_sha256)
    args.report.parent.mkdir(parents=True,exist_ok=True)
    args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"mode":report["mode"],"workflow":report["workflow"],"coverage":{key:value for key,value in report["coverage"].items() if key!="document_hashes"}},ensure_ascii=False))


if __name__=="__main__":
    main()
