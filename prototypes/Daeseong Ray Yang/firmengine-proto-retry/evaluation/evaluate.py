"""Compare actual DOCX bodies/tables and XLSX cells with an independent truth file.

Does not read the generator's content model as expected answers. Truth must be
prepared before running generation; hidden/unreadable values are unscorable.
"""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
from docx import Document
from openpyxl import load_workbook


def normalized(value):
    return " ".join(str(value if value is not None else "").split())


def equivalent(actual, expected):
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        try:
            return Decimal(str(actual)) == Decimal(str(expected))
        except InvalidOperation:
            return False
    return normalized(actual) == normalized(expected)


def evaluate(output_dir: Path, truth_path: Path):
    truth_bytes = truth_path.read_bytes()
    truth = json.loads(truth_bytes)
    results = []
    for rule in truth.get("checks", []):
        result = {"id": rule["id"], "file": rule.get("file"), "type": rule["type"]}
        if rule.get("unscorable"):
            results.append(result | {"status": "unscorable", "reason": rule.get("reason", "독립 기준에서 채점 불가")})
            continue
        path = (output_dir / rule["file"]).resolve()
        if path.parent != output_dir.resolve() or not path.exists():
            results.append(result | {"status": "fail", "reason": "산출물 없음 또는 부적절한 경로"})
            continue
        if rule["type"] == "paragraph_contains":
            body = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
            expected = rule["expected"]
            passed = normalized(expected) in normalized(body)
            result |= {"expected": expected, "actual": "본문에서 일치" if passed else "본문에 없음"}
        elif rule["type"] == "table_rows":
            document = Document(path)
            headers = rule["headers"]
            matches = [table for table in document.tables if [normalized(cell.text) for cell in table.rows[0].cells] == [normalized(h) for h in headers]]
            actual = [[normalized(cell.text) for cell in row.cells] for row in matches[0].rows[1:]] if matches else []
            expected = [[normalized(value) for value in row] for row in rule["expected"]]
            if rule.get("allow_extra_rows"):
                remaining = actual.copy()
                passed = True
                for row in expected:
                    if row in remaining: remaining.remove(row)
                    else: passed = False
            else:
                passed = actual == expected
            result |= {"expected": expected, "actual": actual, "headers": headers}
        elif rule["type"] == "xlsx_cell":
            workbook = load_workbook(path, data_only=False)
            actual = workbook[rule["sheet"]][rule["cell"]].value if rule["sheet"] in workbook.sheetnames else None
            passed = equivalent(actual, rule["expected"])
            result |= {"sheet":rule["sheet"], "cell":rule["cell"], "expected":rule["expected"], "actual":actual}
            workbook.close()
        elif rule["type"] == "manifest_integrity":
            manifest = json.loads(path.read_text(encoding="utf-8"))
            mismatches = []
            for artifact in manifest.get("artifacts", []):
                target = (output_dir / artifact["filename"]).resolve()
                if target.parent != output_dir.resolve() or not target.exists() or hashlib.sha256(target.read_bytes()).hexdigest() != artifact["sha256"]:
                    mismatches.append(artifact["filename"])
            passed = bool(manifest.get("artifacts")) and not mismatches
            result |= {"actual":mismatches,"expected":[]}
        else:
            results.append(result | {"status":"fail","reason":"지원하지 않는 평가 유형"})
            continue
        results.append(result | {"status": "pass" if passed else "fail"})
    return {"evaluated_at": datetime.now(timezone.utc).isoformat(),
            "truth_sha256": hashlib.sha256(truth_bytes).hexdigest(), "metadata":truth.get("metadata", {}),
            "summary": {status:sum(row["status"] == status for row in results) for status in ("pass", "fail", "unscorable")},
            "quality_claim": "합성 문서 생성 경로 확인만 수행; 실제 샘플 제출서류 품질·전문가 정답·시간 단축 미검증" if truth.get("metadata", {}).get("synthetic") else "독립 기준의 채점 결과; 법원 인가율이나 범용 정확도와 구분",
            "results":results}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--truth", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    report = evaluate(args.output_dir, args.truth)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False))
    raise SystemExit(1 if report["summary"]["fail"] else 0)


if __name__ == "__main__":
    main()
