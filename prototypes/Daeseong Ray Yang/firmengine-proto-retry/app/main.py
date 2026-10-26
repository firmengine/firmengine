"""Loopback-only local workflow API. Case data and audit calls stay in ignored .local/."""
from __future__ import annotations

import asyncio
import hashlib
import html
import json
import mimetypes
import os
import shutil
import re
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import fitz
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles

from . import store
from .assessment import assess, synthetic_case

ROOT = store.ROOT
CASE_FIELDS = {'title','applicant_name','address','birth_date','court','filing_date','income_type','tax_type',
               'period_start','period_end','repayment_months','living_cost','deductions','statement','debt_cause',
               'assets_declared','rule_approved','template_approved','family_size'}
COLLECTIONS = {'incomes','creditors','assets','transactions','facts'}
MONEY_FIELDS = {'monthly_gross','monthly_expense','principal','interest','market_value','secured_balance','exempt_amount','amount','living_cost','deductions','sale_price','secured_repayment','costs','replacement_cost'}


def check_input(values):
    for key, value in values.items():
        if key in MONEY_FIELDS and value not in (None,''):
            if isinstance(value, bool) or not isinstance(value, (int,float,str)):
                raise HTTPException(422, f'{key}: 원 단위 정수를 입력하세요.')
            try:
                number = int(value)
                if float(value) != number or number < 0 or number > 9_000_000_000_000:
                    raise ValueError()
            except (ValueError,OverflowError):
                raise HTTPException(422, f'{key}: 0 이상의 원 단위 정수가 필요합니다.')
            values[key] = number
        if key in ('period_start','period_end','filing_date','birth_date','valuation_date','date') and value:
            try:
                date.fromisoformat(value)
            except (TypeError,ValueError):
                raise HTTPException(422, f'{key}: YYYY-MM-DD 날짜가 필요합니다.')
    if values.get('repayment_months') not in (None,''):
        try:
            term = int(values['repayment_months'])
            if term < 1 or term > 120 or float(values['repayment_months']) != term:
                raise ValueError()
            values['repayment_months'] = term
        except (ValueError,TypeError):
            raise HTTPException(422,'변제기간은 1~120 범위의 정수로 입력하고 법적 허용 기간을 별도로 검토하세요.')
    start, end = values.get('period_start'), values.get('period_end')
    if start and end and start > end:
        raise HTTPException(422,'대상기간 시작일이 종료일보다 늦습니다.')
    return values


def expected(body):
    if not isinstance(body.get('expected_revision'), int):
        raise HTTPException(422,'expected_revision 정수가 필요합니다. 최신 사건을 다시 불러오세요.')
    return body['expected_revision']


def object_path(key):
    target = (store.OBJECTS / key).resolve()
    if not target.is_relative_to(store.OBJECTS) or not target.is_file():
        raise HTTPException(404,'파일을 찾을 수 없습니다.')
    return target


async def detail(case_id):
    case = store.get_case(case_id)
    assessment = case.pop('_assessment', None)
    if not assessment or assessment.get('case_revision') != case['revision']:
        assessment = await assess(case)
        store.save_assessment(case_id, case['revision'], assessment)
        latest = store.get_case(case_id)
        if latest['revision'] != case['revision']:
            raise store.Conflict('조회 도중 사건이 변경되었습니다. 다시 불러오세요.')
        case['status'] = latest['status']
    case['assessment'] = assessment
    case['artifacts'] = store.rows('artifacts',case_id)
    for artifact in case['artifacts']:
        artifact.pop('object_key',None)
        if artifact['case_revision'] != case['revision']:
            artifact['previous_status'] = artifact.get('status')
            artifact['status'] = '과거본·재검수 필요'
    case['reviews'] = store.rows('reviews',case_id)
    for review in case['reviews']:
        review['valid'] = review['case_revision'] == case['revision']
    case['events'] = store.rows('events',case_id)
    case['jobs'] = store.rows('jobs',case_id)
    for document in case['documents']:
        document.pop('object_key',None)
        document['url'] = f"/api/cases/{case_id}/documents/{document['id']}/file"
    return case


async def process_job(job):
    case = store.get_case(job['case_id'])
    if case['revision'] != job['revision']:
        store.finish_job(job['id'],'superseded',{'reason':'새 사건 버전 존재'})
        return
    if job['kind'] == 'reassess':
        assessment = await assess(case)
        adopted = store.save_assessment(case['id'],case['revision'],assessment)
        store.finish_job(job['id'],'completed' if adopted else 'superseded',{'can_generate':assessment['can_generate']})
        return
    if job['kind'] == 'analyze':
        from .pipeline import extract_document
        docs = [d for d in case['documents'] if not d.get('deleted') and (d.get('role') == 'evidence' or any(p.get('role') == 'evidence' for p in d.get('pages',[]))) and d.get('status') != 'extracted']
        limit = asyncio.Semaphore(int(os.environ.get('FIRMENGINE_DOCUMENT_CONCURRENCY','2')))

        async def one(document):
            async with limit:
                return document, await extract_document(object_path(document['object_key']), document, case)

        results = await asyncio.gather(*(one(d) for d in docs), return_exceptions=True)
        failures, success = [], []
        for result in results:
            if isinstance(result, Exception):
                failures.append(type(result).__name__)
            else:
                document, output = result
                store.record_calls(case['id'],job['id'],output.get('calls',[]))
                success.append((document, output))
        if not success and failures:
            raise RuntimeError('Document batch failed: '+','.join(failures))

        def adopt(current):
            for document, output in success:
                target = next(d for d in current['documents'] if d['id'] == document['id'])
                logical = output.get('logical_documents',[])
                pending = [p for p in logical if p.get('status') in ('failed','pending','unreadable','manual_review','mapping_review') or p.get('missing_features') or p.get('unreadable_features')]
                target.update(status='needs_review' if pending else 'extracted',warnings=output.get('warnings',[]),logical_documents=logical)
                existing_keys = {(f.get('document_id'),f.get('object_id'),f.get('key'),str(f.get('value'))) for f in current['facts']}
                for fact in output.get('facts',[]):
                    fact = dict(fact)
                    dedup = (document['id'],fact.get('object_id'),fact.get('key'),str(fact.get('value')))
                    if dedup in existing_keys:
                        continue
                    fact.update(id=store.uid(),document_id=document['id'])
                    fact['status'] = 'candidate' # machine candidates cannot approve legal recognition
                    # Verified nonlegal identity text may fill empty consultation fields;
                    # conflicts remain visible rather than overwriting the current user value.
                    key = fact.get('key')
                    if key in ('applicant_name','address','birth_date') and fact.get('verification',{}).get('quote_verified'):
                        value = str(fact.get('value',''))
                        if not current.get(key):
                            if key != 'birth_date' or re.fullmatch(r'\d{4}-\d{2}-\d{2}',value):
                                current[key] = value
                                fact['status'] = 'confirmed'
                                fact['reason'] = '원문 문자층 인용 대응 검증 후 빈 상담 항목에 자동 반영; 법적 인정값과 구분'
                        elif str(current[key]).strip() != value.strip():
                            fact['status'] = 'disputed'
                            fact['reason'] = '현재 상담 값과 증빙 인용의 값이 달라 원본 대조 필요'
                    current['facts'].append(fact)
                for collection, entities in (output.get('entities') or {}).items():
                    if collection not in COLLECTIONS:
                        continue
                    for entity in entities:
                        row = dict(entity)
                        source_object = row.get('source_object_id')
                        if source_object and any(r.get('source_object_id')==source_object for r in current[collection]):
                            continue
                        row.update(id=store.uid(),confirmed=False,source_document_id=document['id'])
                        current[collection].append(row)
            current['analysis_warnings'] = failures
        updated = store.mutate(case['id'],job['revision'],adopt,'documents.extracted','증빙 전용 추출 결과 채택',preserve_job_id=job['id'])
        logical_results = [page for _, output in success for page in output.get('logical_documents',[])]
        evidence_pages = [page for page in logical_results if page.get('role') == 'evidence']
        failed_pages = [page for page in evidence_pages if page.get('status') == 'failed']
        waiting_pages = [page for page in evidence_pages if page.get('status') in ('pending','unreadable','manual_review','mapping_review')]
        outcome = 'failed' if evidence_pages and len(failed_pages)==len(evidence_pages) else 'partial' if failures or failed_pages or waiting_pages else 'completed'
        store.finish_job(job['id'],outcome,{'documents':len(success),'failures':failures,'failed_pages':len(failed_pages),'pending_pages':len(waiting_pages),'adopted_revision':updated['revision']})


async def worker(stop):
    while not stop.is_set():
        job = store.claim_job()
        if not job:
            await asyncio.sleep(.4)
            continue
        try:
            await process_job(job)
        except KeyError:
            store.purge_case_cache(job['case_id'])
            store.finish_job(job['id'],'superseded',{'reason':'삭제된 사건의 지연 완료 결과 폐기·캐시 정리'})
        except store.Conflict:
            store.finish_job(job['id'],'superseded',{'reason':'사건 삭제 또는 새 버전'})
        except asyncio.CancelledError:
            store.finish_job(job['id'],'queued',{'reason':'서버 종료 후 재개 대기'})
            raise
        except Exception as exc:
            retry = job['attempts'] < 3
            run_at = (datetime.now(timezone.utc)+timedelta(seconds=3*job['attempts'])).isoformat(timespec='seconds')
            store.finish_job(job['id'],'retry' if retry else 'failed',{'error_type':type(exc).__name__,'message':'처리 실패: 원본과 런타임 연결을 확인하세요.'},run_at if retry else None)


@asynccontextmanager
async def lifespan(app):
    store.initialize()
    stop = asyncio.Event()
    task = asyncio.create_task(worker(stop))
    yield
    stop.set()
    task.cancel()
    try:
        await task
    except asyncio.CancelledError:
        pass


app = FastAPI(title='Firm Engine · Retry Prototype',version='0.1.0',lifespan=lifespan)


@app.middleware('http')
async def local_boundary(request: Request, call_next):
    # No cloud exposure or cross-origin case writes. No case content in HTTP logging.
    host = request.headers.get('host','').split(':')[0]
    if host not in ('127.0.0.1','localhost','[','testserver'):
        return JSONResponse({'detail':'로컬 주소에서만 접속할 수 있습니다.'},status_code=403)
    origin = request.headers.get('origin')
    if origin and request.method not in ('GET','HEAD','OPTIONS'):
        from urllib.parse import urlparse
        parsed = urlparse(origin)
        if parsed.netloc != request.headers.get('host'):
            return JSONResponse({'detail':'다른 출처의 변경 요청을 허용하지 않습니다.'},status_code=403)
    response = await call_next(request)
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    return response


@app.exception_handler(store.Conflict)
async def conflict_handler(request,exc):
    return JSONResponse({'detail':str(exc)},status_code=409)


@app.exception_handler(KeyError)
async def missing_handler(request,exc):
    return JSONResponse({'detail':'대상을 찾을 수 없습니다.'},status_code=404)


@app.get('/api/status')
async def status():
    from .ai import status as ai_status
    from .pipeline import runtime_status
    from .rag import status as vector_status
    return dict(ai=ai_status(),n8n=await runtime_status(),storage='SQLite · 로컬 영속 원장',prototype=True,
                auth='로컬 검수자 이름 입력 · 운영용 인증 미구현',external_sync=False,vector_search=vector_status())


@app.get('/api/cases')
async def cases():
    keys = ('id','case_number','title','applicant_name','court','income_type','revision','status','updated_at','archived','is_demo')
    return {'cases':[{key:case.get(key) for key in keys} for case in store.list_cases()]}


@app.post('/api/cases',status_code=201)
async def create(body:dict):
    values = check_input({k:v for k,v in body.items() if k in CASE_FIELDS})
    defaults = store.get_setting('reviewed_release',{})
    for k in ('rule_approved','template_approved'):
        if k not in values:
            values[k] = defaults.get(k,False)
    case = store.create_case(values)
    return await detail(case['id'])


@app.post('/api/demo',status_code=201)
async def demo():
    case = store.create_case(synthetic_case())
    return await detail(case['id'])


@app.get('/api/cases/{case_id}')
async def read_case(case_id:str):
    return await detail(case_id)


@app.patch('/api/cases/{case_id}')
async def edit_case(case_id:str,body:dict):
    revision = expected(body)
    changes = check_input({k:v for k,v in body.items() if k in CASE_FIELDS})
    candidate = {**store.get_case(case_id),**changes}
    check_input(candidate)
    store.mutate(case_id,revision,lambda c:c.update(changes),'case.updated',body.get('reason') or '상담정보 정정')
    return await detail(case_id)


@app.post('/api/cases/{case_id}/archive')
async def archive(case_id:str,body:dict):
    store.mutate(case_id,expected(body),lambda c:c.update(archived=not c.get('archived')),'case.archive_changed',body.get('reason','사건 보관 상태 변경'))
    return await detail(case_id)


@app.delete('/api/cases/{case_id}')
async def remove_case(case_id:str):
    store.delete_case(case_id)
    return {'deleted':True,'policy':'사건·원본·생성물·호출기록·검수이력 삭제. 사용자가 내려받은 외부 파일은 별도 보관.'}


def validate_entity(case_id,collection,values):
    if collection not in COLLECTIONS:
        raise HTTPException(404,'지원하지 않는 항목 종류입니다.')
    values = check_input(values)
    if collection == 'facts' and values.get('status') == 'confirmed' and not (values.get('reason') or values.get('quote')):
        raise HTTPException(422,'사실 확정 시 원문 근거 또는 확인 사유를 남기세요.')
    source_id = values.get('source_document_id') or values.get('document_id')
    if source_id:
        docs = store.get_case(case_id)['documents']
        source = next((d for d in docs if d['id']==source_id and not d.get('deleted')),None)
        if not source or source.get('role') != 'evidence':
            raise HTTPException(422,'의뢰인 증빙으로 매핑한 원본에만 사실을 연결할 수 있습니다.')
        source_page = values.get('source_page') or values.get('page')
        if source_page:
            valid_pages = {int(p.get('original_page',p['page'])) for p in source.get('pages',[])} or set(range(1,source['page_count']+1))
            if int(source_page) not in valid_pages:
                raise HTTPException(422,'이 증빙에 포함된 원본 페이지가 아닙니다.')
    return values


def apply_fact_to_subject(case, fact):
    """Correct source facts; derived calculations are always recomputed."""
    if fact.get('status') != 'confirmed':
        return
    key, value = fact.get('key'),fact.get('value')
    if key in CASE_FIELDS:
        normalized = check_input({key:value})
        case.update(normalized)
    subject = fact.get('object_id')
    aliases = {'creditor_name':'name','business_name':'label','employer':'label','asset_name':'label','withdrawal_amount':'amount','deposit_amount':'amount','transaction_date':'date','description':'purpose'}
    field = aliases.get(key,key)
    allowed = {'monthly_gross','monthly_expense','period_start','period_end','name','label','contract_id','principal','interest','security','priority','market_value','secured_balance','valuation_date','valuation_source','amount','date','purpose','counterparty'}
    if subject and field in allowed:
        for collection in ('incomes','creditors','assets','transactions'):
            for row in case[collection]:
                if row.get('source_object_id') == subject:
                    row.update(check_input({field:value}))
                    # Correcting a numeric claim does not silently approve its legal treatment.
                    if collection in ('incomes','creditors','assets'):
                        row['confirmed'] = False


@app.post('/api/cases/{case_id}/entities/{collection}')
async def add_entity(case_id:str,collection:str,body:dict):
    values = validate_entity(case_id,collection,{k:v for k,v in body.items() if k not in ('expected_revision',)})
    values['id'] = store.uid()
    def add(c):
        c[collection].append(values)
        if collection == 'facts':
            apply_fact_to_subject(c,values)
    store.mutate(case_id,expected(body),add,collection+'.created',body.get('reason') or '항목 추가')
    return await detail(case_id)


@app.patch('/api/cases/{case_id}/entities/{collection}/{entity_id}')
async def edit_entity(case_id:str,collection:str,entity_id:str,body:dict):
    if collection not in COLLECTIONS:
        raise HTTPException(404)
    case = store.get_case(case_id)
    old = next((r for r in case[collection] if r['id']==entity_id),None)
    if old is None:
        raise HTTPException(404,'항목이 없습니다.')
    values = validate_entity(case_id,collection,{**old,**{k:v for k,v in body.items() if k not in ('expected_revision','id')}})
    def update(c):
        c[collection] = [values if r['id']==entity_id else r for r in c[collection]]
        if collection == 'facts':
            apply_fact_to_subject(c,values)
    store.mutate(case_id,expected(body),update,collection+'.updated',body.get('reason') or '항목 정정')
    return await detail(case_id)


@app.delete('/api/cases/{case_id}/entities/{collection}/{entity_id}')
async def remove_entity(case_id:str,collection:str,entity_id:str,expected_revision:int):
    if collection not in COLLECTIONS:
        raise HTTPException(404)
    def remove(c):
        if not any(r['id']==entity_id for r in c[collection]):
            raise KeyError(entity_id)
        c[collection] = [r for r in c[collection] if r['id']!=entity_id]
    store.mutate(case_id,expected_revision,remove,collection+'.removed','항목 제외; 이전 값은 이력 보존')
    return await detail(case_id)


@app.post('/api/cases/{case_id}/documents',status_code=201)
async def upload(case_id:str,file:UploadFile=File(...),role:str=Form('unknown'),document_type:str=Form('unknown'),
                 period_start:str=Form(''),period_end:str=Form(''),expected_revision:int|None=Form(None)):
    case = store.get_case(case_id)
    if role not in ('evidence','evaluation','administrative','unknown'):
        raise HTTPException(422,'자료 역할을 지정하세요.')
    suffix = Path(file.filename or '').suffix.lower()
    if suffix not in ('.pdf','.jpg','.jpeg','.png'):
        raise HTTPException(422,'PDF·JPG·PNG를 업로드하세요.')
    data = await file.read(60*1024*1024+1)
    if len(data)>60*1024*1024:
        raise HTTPException(413,'현재 로컬 실행 한도는 파일당 60MB입니다.')
    try:
        with fitz.open(stream=data,filetype=suffix.lstrip('.')) as pdf:
            pages = len(pdf)
    except Exception:
        raise HTTPException(422,'파일을 읽을 수 없습니다. 정상 원본을 다시 제출하세요.')
    identifier = store.uid()
    key = f'originals/{identifier}{suffix}'
    path = store.OBJECTS / key
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_bytes(data)
    doc = dict(id=identifier,name=Path(file.filename or 'document').name,role=role,document_type=document_type,
               period_start=period_start,period_end=period_end,page_count=pages,object_key=key,
               sha256=hashlib.sha256(data).hexdigest(),storage_provider='local',namespace='originals',status='uploaded',
               pages=[dict(page=n,role=role,document_type=document_type) for n in range(1,pages+1)],created_at=store.now())
    try:
        store.mutate(case_id,expected_revision if expected_revision is not None else case['revision'],lambda c:c['documents'].append(doc),'document.uploaded','원본 업로드 · 역할 '+role)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return await detail(case_id)


@app.patch('/api/cases/{case_id}/documents/{document_id}')
async def map_document(case_id:str,document_id:str,body:dict):
    def update(c):
        document = next(d for d in c['documents'] if d['id']==document_id and not d.get('deleted'))
        for key in ('role','document_type','period_start','period_end'):
            if key in body:
                document[key] = body[key]
        if document['role'] not in ('evidence','evaluation','administrative','unknown'):
            raise HTTPException(422,'잘못된 자료 역할입니다.')
        if 'pages' in body:
            supplied = body['pages']
            if sorted(int(p['page']) for p in supplied) != list(range(1,document['page_count']+1)):
                raise HTTPException(422,'모든 원본 페이지를 중복 없이 매핑하세요.')
            if any(p.get('role') not in ('evidence','evaluation','administrative','unknown') for p in supplied):
                raise HTTPException(422,'페이지별 자료 역할이 올바르지 않습니다.')
            document['pages'] = supplied
        elif 'role' in body or 'document_type' in body:
            document['pages'] = [dict(page=n,role=document['role'],document_type=document['document_type']) for n in range(1,document['page_count']+1)]
        document['status'] = 'uploaded'
        # Old machine candidates from re-mapped material cannot silently stay usable.
        c['facts'] = [f for f in c['facts'] if f.get('document_id') != document_id]
        for collection in ('incomes','creditors','assets'):
            for row in c[collection]:
                if row.get('source_document_id') == document_id:
                    row['confirmed'] = False
    store.mutate(case_id,expected(body),update,'document.mapped',body.get('reason') or '문서 역할·페이지·기간 정정')
    return await detail(case_id)


@app.delete('/api/cases/{case_id}/documents/{document_id}')
async def remove_document(case_id:str,document_id:str,expected_revision:int):
    def remove(c):
        row = next(d for d in c['documents'] if d['id']==document_id)
        row.update(deleted=True,status='excluded',deleted_at=store.now())
        for collection in ('incomes','creditors','assets'):
            for value in c[collection]:
                if value.get('source_document_id') == document_id:
                    value['confirmed'] = False
        c['facts'] = [f for f in c['facts'] if f.get('document_id') != document_id]
    store.mutate(case_id,expected_revision,remove,'document.excluded','자료 제외; 원본은 사건 삭제 시 함께 삭제')
    return await detail(case_id)


@app.get('/api/cases/{case_id}/documents/{document_id}/file')
async def original_file(case_id:str,document_id:str):
    row = next((d for d in store.get_case(case_id)['documents'] if d['id']==document_id and not d.get('deleted')),None)
    if not row:
        raise HTTPException(404)
    return FileResponse(object_path(row['object_key']),filename=row['name'])


@app.post('/api/cases/{case_id}/analyze')
async def analyze(case_id:str,body:dict):
    job_id = store.enqueue(case_id,expected(body),'analyze')
    result = await detail(case_id)
    result['job_id'] = job_id
    return result


@app.post('/api/cases/{case_id}/generate')
async def generate(case_id:str,body:dict):
    from .documents import generate_package
    case = store.get_case(case_id)
    if expected(body) != case['revision']:
        raise store.Conflict('최신 사건 버전에서 다시 생성하세요.')
    assessment = await assess(case)
    store.save_assessment(case_id,case['revision'],assessment)
    if not assessment['can_generate']:
        raise HTTPException(409,dict(message='필수 조건이 남아 법원 서류 생성을 보류합니다.',blockers=assessment['blockers']))
    out = store.OBJECTS / 'generated' / store.uid()
    out.mkdir(parents=True)
    try:
        generated = await asyncio.to_thread(generate_package,case,assessment,out)
        for item in generated:
            item['rule_version'] = assessment.get('rule_version')
            item['calculation_code_sha256'] = assessment['calculations'].get('code_sha256')
        store.add_artifacts(case_id,case['revision'],generated)
    except Exception:
        # Concrete, checked owned directory only. Discard stale unfinished generation.
        if out.resolve().is_relative_to(store.OBJECTS.resolve()):
            shutil.rmtree(out)
        raise
    return await detail(case_id)


@app.get('/api/artifacts/{artifact_id}/file')
async def artifact_file(artifact_id:str):
    row = store.artifact(artifact_id)
    path = object_path(row['object_key'])
    if hashlib.sha256(path.read_bytes()).hexdigest() != row['sha256']:
        raise HTTPException(409,'생성 파일의 해시가 일치하지 않습니다.')
    return FileResponse(path,filename=path.name)


@app.post('/api/cases/{case_id}/reviews')
async def review(case_id:str,body:dict):
    if body.get('role') not in ('staff','lawyer') or not str(body.get('reviewer','')).strip():
        raise HTTPException(422,'검수 역할과 검수자 이름을 입력하세요.')
    current = await detail(case_id)
    if not current['assessment']['can_generate']:
        raise HTTPException(409,'필수 조건을 보완하고 다시 생성하세요.')
    store.add_review(case_id,body['artifact_id'],body['role'],body['reviewer'],expected(body))
    return await detail(case_id)


@app.get('/api/cases/{case_id}/review-report')
async def review_report(case_id:str):
    from docx import Document
    case = await detail(case_id)
    out = store.OBJECTS / 'reports' / (store.uid()+'.docx')
    out.parent.mkdir(exist_ok=True)
    doc = Document()
    doc.add_heading('Firm Engine · 생성본 검수 기록',0)
    doc.add_paragraph(f"사건: {case['title']} / 버전 {case['revision']} / 출력 {store.now()}")
    doc.add_paragraph('로컬 프로토타입 검수 기록. 운영용 인증이나 법원 접수·인가를 의미하지 않습니다.')
    for item in case['artifacts']:
        doc.add_heading(item['title'],2)
        doc.add_paragraph(f"생성본 {item['id']} / 사건 v{item['case_revision']} / {item['status']}\nSHA-256 {item['sha256']}")
        for record in case['reviews']:
            if record['artifact_id']==item['id']:
                doc.add_paragraph(f"{record['role']} · {record['reviewer']} · {record['created_at']} · {'유효' if record['valid'] else '과거본 확인'}")
    doc.save(out)
    # The response is transient; no unattached report is retained beyond the download.
    from starlette.background import BackgroundTask
    return FileResponse(out,filename='생성본-검수기록.docx',background=BackgroundTask(out.unlink,missing_ok=True))


@app.post('/api/cases/{case_id}/scenario')
async def scenario(case_id:str,body:dict):
    values = check_input(dict(body))
    case = store.get_case(case_id)
    row = next((a for a in case['assets'] if a['id']==body.get('asset_id')),None)
    if not row:
        raise HTTPException(404,'대상 재산을 선택하세요.')
    sale, secured, costs, replacement = [int(values.get(k) or 0) for k in ('sale_price','secured_repayment','costs','replacement_cost')]
    net = sale-secured-costs
    return dict(asset_id=row['id'],net_cash=net-replacement,replacement_asset=replacement,total_after=net,
                original_market_value=row.get('market_value'),scenario_only=True,
                explanation='매각대금 - 담보상환 - 비용 = 잔여 재산. 대체재산 취득은 현금과 대체재산 사이 이동이며 재산 총액에서 중복 차감하지 않습니다. 실제 처분·법적 청산가치 인정은 증빙과 별도 검토가 필요합니다.')


@app.get('/api/knowledge')
async def knowledge():
    target = ROOT/'knowledge'/'catalog.json'
    data = json.loads(target.read_text(encoding='utf-8')) if target.exists() else {'sources':[]}
    from .rag import status as vector_status
    data['capabilities'] = {'vector_search':vector_status(),'success_case_collection':'사용자 검토 자료만 적재; 평가 사건 그룹 제외','scheduled_law_update':'워크플로 정의·수동 활성화 필요','remote_valuation':'인증 공급자 미연결; 평가 출처 직접 입력'}
    return data


@app.get('/api/knowledge/vector-status')
async def vector_status():
    from .rag import status
    return status()


@app.post('/api/knowledge/search')
async def vector_search(body:dict):
    from .rag import search
    try:
        filters = dict(body.get('filters') or body)
        filters.pop('query',None)
        return await search(body.get('query',''),filters)
    except ValueError as exc:
        raise HTTPException(422,str(exc))


@app.post('/api/knowledge/ingest')
async def vector_ingest(body:dict):
    from .rag import ingest
    try:
        return await ingest(body.get('text',''),body.get('metadata',{}))
    except ValueError as exc:
        raise HTTPException(422,str(exc))


@app.post('/api/knowledge/refresh')
async def refresh_knowledge():
    """Collect official source editions only; raw text changes never approve a rule."""
    import httpx
    from urllib.parse import urlparse
    catalog = await knowledge()
    results = []
    async with httpx.AsyncClient(timeout=25,trust_env=False,follow_redirects=False) as client:
        for source in catalog.get('sources',[]):
            url = source.get('url')
            if not url or urlparse(url).hostname not in ('www.law.go.kr','law.go.kr','slb.scourt.go.kr'):
                continue
            key = 'official-source:'+source['id']
            prior = store.get_setting(key,{})
            try:
                response = await client.get(url)
                response.raise_for_status()
                digest = hashlib.sha256(response.content).hexdigest()
                snapshot = store.LOCAL/'knowledge'/source['id']/f'{digest}.source'
                snapshot.parent.mkdir(parents=True,exist_ok=True)
                snapshot.write_bytes(response.content)
                record = dict(source_id=source['id'],url=url,sha256=digest,previous_sha256=prior.get('sha256'),retrieved_at=store.now(),
                              changed=bool(prior.get('sha256') and prior['sha256']!=digest),status='원문 수집·의미 검토 필요',
                              automatic_semantic_approval=False,version=source.get('source_version'),effective_date=source.get('effective_date'))
                store.set_setting(key,record)
                results.append(record)
            except Exception as exc:
                results.append(dict(source_id=source['id'],status='조회 실패',error_type=type(exc).__name__,next_action='공식 원문을 수동 확인하세요.'))
    store.set_setting('last_law_refresh',dict(created_at=store.now(),results=results))
    return dict(results=results,scheduled_firing_verified=False,message='수집 성공은 규칙 승인이나 모든 법령 변경의 자동 적용을 의미하지 않습니다.')


@app.get('/api/promptbook')
async def promptbook():
    from .ai import list_prompts
    return list_prompts()


@app.put('/api/promptbook/{name}')
async def edit_prompt(name:str,body:dict):
    from .ai import update_prompt
    try:
        return update_prompt(name,{k:v for k,v in body.items() if k in ('system','body')})
    except ValueError as exc:
        raise HTTPException(422,str(exc))


@app.get('/api/jobs')
async def jobs():
    return {'jobs':store.rows('jobs')}


@app.get('/api/cases/{case_id}/calls')
async def calls(case_id:str):
    store.get_case(case_id)
    return {'calls':store.rows('calls',case_id)}


@app.post('/api/approvals')
async def release_approval(body:dict):
    if not body.get('reviewer') or not body.get('reason'):
        raise HTTPException(422,'판본 검토자와 사유가 필요합니다.')
    value = {k:body.get(k,False) for k in ('rule_approved','template_approved')}
    value.update(reviewer=body['reviewer'],reason=body['reason'],version='retry-v1',created_at=store.now())
    store.set_setting('reviewed_release',value)
    return value


@app.get('/')
async def index():
    return FileResponse(ROOT/'static'/'index.html')


app.mount('/static',StaticFiles(directory=ROOT/'static',check_dir=False),name='static')
