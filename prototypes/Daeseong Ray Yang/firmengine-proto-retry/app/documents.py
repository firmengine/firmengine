"""Editable, evidence-linked court-form drafts built from one preview/file model.

Fresh implementation against the current spec. Public HWP originals are retained in
templates/originals; the DOCX mapping is a reconstruction awaiting expert validation.
No submitted sample, evaluation answer, or earlier prototype is used here.
"""
from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from html import escape
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_VERSION = "retry-court-mapping-2026-10-26.1"
RENDERER_VERSION = "retry-renderer-2026-10-26.3-empty-sum-unknown"
UNKNOWN = "확인 필요"
DIAGNOSTIC_BANNER = "진단 전용 · 미확정 추출 후보 · 실제 제출서류 생성 성공 아님"
FORM_IDS = ("D5100", "D5106", "D5101", "D5103", "D5105", "D5110")


def _safe(value):
    if value is None or value == "":
        return UNKNOWN
    if isinstance(value, bool):
        return "예" if value else "아니오"
    return str(value)


def _number(value):
    if value is None or value == "" or isinstance(value, bool):
        return None
    try:
        number = Decimal(str(value).replace(",", ""))
        return int(number) if number.is_finite() else None
    except (InvalidOperation, ValueError, OverflowError):
        return None


def _money(value):
    amount = _number(value)
    return f"{amount:,}원" if amount is not None else UNKNOWN


def _field(case, key, default=None):
    value = case.get(key)
    if value is not None and value != "":
        return value
    facts = [f for f in case.get("facts", []) if f.get("key") == key and f.get("status") == "confirmed"]
    if facts:
        return facts[-1].get("value")
    if case.get("diagnostic_mode"):
        candidates = [fact.get("value") for fact in case.get("facts", [])
                      if fact.get("key") == key and fact.get("status") == "candidate"
                      and fact.get("value") is not None and fact.get("value") != ""]
        normalized = {json.dumps(value, ensure_ascii=False, sort_keys=True) for value in candidates}
        if len(normalized) == 1:
            return candidates[0]
    return default


def _total(rows, key):
    if not rows:
        return None
    values = [_number(row.get(key)) for row in rows]
    return sum(values) if all(value is not None for value in values) else None


def _p(text):
    return {"type": "paragraph", "text": _safe(text)}


def _heading(text):
    return {"type": "heading", "text": text}


def _table(headers, rows, name=None):
    return {"type": "table", "headers": headers, "rows": [[_safe(cell) for cell in row] for row in rows], "name": name}


def _identity(case):
    return _table(["항목", "기재 내용"], [
        ["신청인", _field(case, "applicant_name")],
        ["주민등록번호", _field(case, "resident_number")],
        ["생년월일", _field(case, "birth_date")],
        ["주민등록상 주소", _field(case, "registered_address", _field(case, "address"))],
        ["현재 주소", _field(case, "address")],
        ["송달 장소", _field(case, "service_address")],
        ["휴대전화", _field(case, "phone")],
        ["대리인 / 사무소", _field(case, "representative")],
        ["관련 절차 사건번호", _field(case, "related_case_number")],
        ["반환받을 본인 계좌", _field(case, "refund_account")],
    ], "신청인 인적사항")


def _signature(case):
    return [_p(f"신청일: {_safe(case.get('filing_date'))}"),
            _p(f"신청인 {_safe(case.get('applicant_name'))}  (서명 또는 날인: ____________________)"),
            _p(f"{_safe(case.get('court'))} 귀중")]


def _allocations(case, calc):
    creditors = {str(row.get("id")): row for row in case.get("creditors", [])}
    rows = []
    for index, row in enumerate(calc.get("allocations", []), 1):
        ref = creditors.get(str(row.get("creditor_id", row.get("id"))), {})
        principal = row.get("principal", ref.get("principal"))
        monthly = row.get("monthly_payment", row.get("monthly", row.get("monthly_amount")))
        total = row.get("total_payment", row.get("total", row.get("total_amount")))
        rows.append([index, row.get("name", row.get("creditor", ref.get("name"))), ref.get("contract_id"),
                     _money(principal), _money(monthly), _money(total),
                     "유보" if row.get("reserved", ref.get("disputed")) else "확정채권"])
    if not rows:
        rows = [["—", "검증된 배분 결과 없음", UNKNOWN, UNKNOWN, UNKNOWN, UNKNOWN, "계산 조건 확인"]]
    return _table(["번호", "채권자", "계약 식별자", "기준 채권액", "월 변제예정액", "총 변제예정액", "확정 / 유보"], rows, "개인회생채권 변제예정액표")


def content_models(case: dict, assessment: dict) -> list[dict]:
    """The authoritative text/table model used by DOCX, previews and evaluation."""
    calc = assessment.get("calculations") or {}
    creditors = case.get("creditors", [])
    assets = case.get("assets", [])
    incomes = case.get("incomes", [])
    models = []

    def add(code, title, blocks, source=None):
        intro = [_p("가상 사건 · 기능 시연용 자료 · 실제 제출용 아님")] if case.get("is_demo") else []
        if case.get("diagnostic_mode"):
            intro = [_p(DIAGNOSTIC_BANNER)] + intro
        models.append({"code": code, "title": title, "source_form": source or code,
                       "blocks": intro + blocks, "template_version": TEMPLATE_VERSION})

    add("D5100", "개인회생절차 개시신청서", [
        _identity(case), _heading("신청 취지"),
        _p("신청인에 대한 개인회생절차의 개시를 구합니다."),
        _heading("신청 이유"),
        _p(f"신청인이 현재 부담하는 채무는 별지 개인회생채권자목록과 같습니다. 채무 원금 합계는 {_money(calc.get('total_principal'))}, 이자 합계는 {_money(calc.get('total_interest'))}입니다."),
        _p(f"현재 수입·지출과 재산은 별지 목록과 같습니다. 소득 유형은 {_safe(case.get('income_type'))}이며, 확인된 월 순소득은 {_money(calc.get('net_income'))}입니다."),
        _p(f"채무 발생 및 신청에 이르게 된 사정: {_safe(case.get('debt_cause'))}"),
        _p(f"변제계획상의 예정 기간은 {_safe(case.get('repayment_months'))}개월, 월 변제예정액은 {_money(calc.get('monthly_payment'))}, 명목 총변제예정액은 {_money(calc.get('total_payment'))}입니다."),
        _p(f"최초 납입예정일: {_safe(_field(case, 'first_payment_date'))}. 매월 납입일: {_safe(_field(case, 'payment_day'))}. 이 날짜는 자료에 근거하여 보완합니다."),
        _heading("첨부서류"),
        _table(["서류", "편철 / 확인"], [["개인회생채권자목록", "생성본 참조"], ["재산목록", "생성본 참조"], ["수입 및 지출에 관한 목록", "생성본 참조"], ["진술서", "생성본 참조"], ["소득·채무·재산 소명자료", "별도 증빙목록 참조"], ["비용 납부·위임장·본인 계좌 사본", "실제 제출 전 확인"]]),
    ] + _signature(case))

    debt_rows = []
    for index, row in enumerate(creditors, 1):
        debt_rows.append([index, row.get("name"), row.get("contract_id"), row.get("cause", row.get("reason")),
                          row.get("address"), _money(row.get("principal")), _money(row.get("interest")),
                          "담보" if row.get("security") == "secured" else "무담보",
                          "다툼 있음" if row.get("disputed") else ("우선권 검토" if row.get("priority") == "tax" else "일반")])
    add("D5106", "개인회생채권자목록", [
        _p(f"채무자: {_safe(case.get('applicant_name'))} / 채권현재액 기준일: {_safe(case.get('filing_date'))}"),
        _table(["번호", "채권자", "계약 식별자", "채권 원인", "주소", "원금", "이자", "담보", "권리 / 다툼"], debt_rows or [[UNKNOWN]*9], "채권자별 채권현재액"),
        _table(["구분", "금액"], [["원금 합계", _money(calc.get("total_principal"))], ["이자 합계", _money(calc.get("total_interest"))], ["채권현재액 합계", _money(calc.get("total_debt"))]]),
        _p("동일 금융기관의 복수 계약은 계약 식별자별로 구분하였습니다. 담보·우선권·다툼은 각각 별도 축으로 검토합니다."),
        _p("채권자별 산정근거 및 부속서류는 증빙목록·조건부 부속표를 함께 대조하여 확인합니다."),
    ])

    asset_rows = []
    for index, row in enumerate(assets, 1):
        market, secured, exempt = [_number(row.get(key)) for key in ("market_value", "secured_balance", "exempt_amount")]
        net = max(0, market-secured) if market is not None and secured is not None else None
        liquid = max(0, net-exempt) if net is not None and exempt is not None else None
        asset_rows.append([index, row.get("label"), row.get("kind"), _money(market), _money(secured), _money(net), _money(exempt), _money(liquid),
                           f"{_safe(row.get('valuation_source'))} / {_safe(row.get('valuation_date'))}"])
    if not assets and case.get("assets_declared"):
        asset_rows = [["—", "등록 재산 없음으로 진술", "담당자 확인", "0원", "0원", "0원", "0원", "0원", "재산 조회 증빙 확인"]]
    add("D5101", "재산목록", [
        _p(f"신청인: {_safe(case.get('applicant_name'))} / 평가 기준일: {_safe(case.get('filing_date'))}"),
        _table(["번호", "명칭", "종류", "시가", "담보잔액", "순가치", "면제 인정액", "청산가치", "평가 출처 / 기준일"], asset_rows or [[UNKNOWN]*9], "재산별 내역"),
        _p(f"청산가치 합계: {_money(calc.get('liquidation_value'))}"),
        _p("시가·순가치·법적 청산가치는 구분하여 기재합니다. 면제 인정액은 담당자가 근거를 확인한 입력값이며 실제 법원의 결정을 뜻하지 않습니다."),
        _p(f"압류·가압류 및 권리관계: {_safe(_field(case, 'asset_seizure_details'))}"),
    ])

    income_rows = []
    for row in incomes:
        gross, expense = _number(row.get("monthly_gross")), _number(row.get("monthly_expense"))
        net = gross-expense if gross is not None and expense is not None else None
        income_rows.append([row.get("label"), row.get("kind"), f"{_safe(row.get('period_start'))} ~ {_safe(row.get('period_end'))}",
                            _money(gross), _money(expense), _money(net), _money(net*12 if net is not None else None),
                            "확인" if row.get("confirmed") else "미확인"])
    family = case.get("family", [])
    add("D5103", "수입 및 지출에 관한 목록", [
        _heading("I. 현재의 수입 목록"),
        _table(["사업체 / 직장 / 소득원", "소득 유형", "산정 대상기간", "월 총수입", "월 인정경비", "월 순수입", "월 순수입 연환산", "확인"], income_rows or [[UNKNOWN]*8], "소득원별 수입"),
        _p(f"월 평균 순수입 합계: {_money(calc.get('net_income'))}. 연환산은 확정 월 순수입 × 12의 표시이며 과거 실제 연소득 합계와 구분합니다."),
        _heading("II. 변제계획 수행 시 예상 지출"),
        _table(["항목", "월 금액", "확인 내용"], [["인정 생계비", _money(case.get("living_cost")), _field(case, "living_cost_reason")],
                                                    ["별도 공제", _money(case.get("deductions")), _field(case, "deductions_reason")],
                                                    ["가용소득", _money(calc.get("disposable_income")), "확정 순소득 - 인정 생계비 - 공제"]]),
        _heading("III. 가족관계"),
        _table(["관계", "성명", "연령", "동거", "직업", "월 수입", "재산", "부양 여부"],
               [[r.get("relationship"), r.get("name"), r.get("age"), r.get("cohabiting"), r.get("occupation"), _money(r.get("income")), _money(r.get("assets")), r.get("dependent")] for r in family] or [[UNKNOWN]*8], "가족관계"),
        _p(f"생계비 산정 가구원 수: {_safe(_field(case, 'family_size'))}. 부양 인정과 추가 지출 사유는 관련 증빙에 따라 별도 검토합니다."),
    ])

    add("D5105", "진술서", [
        _p(f"진술인: {_safe(case.get('applicant_name'))}"),
        _heading("I. 경력"),
        _p(f"최종 학력: {_safe(_field(case, 'education'))}"),
        _table(["기간", "업종 / 소득 유형", "직장 / 사업체", "직위 / 업무"],
               [[f"{_safe(r.get('period_start'))} ~ {_safe(r.get('period_end'))}", r.get("kind"), r.get("label"), r.get("position")] for r in incomes] or [[UNKNOWN]*4]),
        _p(f"결혼·이혼 이력: {_safe(_field(case, 'marital_history'))}"),
        _heading("II. 현재 주거 상황"),
        _p(f"주소: {_safe(case.get('address'))}"),
        _p(f"주거 형태 및 거주 시작일: {_safe(_field(case, 'housing_details'))}"),
        _heading("III. 부채 상황"),
        _p(f"소송·지급명령·전부명령·압류 등: {_safe(_field(case, 'enforcement_details'))}"),
        _p(f"채무를 부담하게 된 사정: {_safe(case.get('debt_cause'))}"),
        _p(f"개인회생 신청에 이르게 된 구체적 사정: {_safe(case.get('statement'))}"),
        _heading("IV. 과거 절차 및 기타 확인"),
        _p(f"과거 회생·파산 등 절차: {_safe(_field(case, 'past_proceedings'))}"),
        _p("이 진술서의 서술은 상담·증빙에서 확인하여 입력한 사실에 기초합니다. 확인되지 않은 사실은 확인 항목으로 남겨두었습니다."),
    ] + _signature(case))

    special = [r for r in creditors if r.get("security") == "secured" or r.get("priority") == "tax" or r.get("disputed") or r.get("estate_claim") or r.get("assignment_order")]
    add("D5110", "변제계획안", [
        _p(f"사건: {_safe(case.get('case_number'))} 개인회생 / 채무자: {_safe(case.get('applicant_name'))}"),
        _heading("1. 변제기간"),
        _p(f"{_safe(_field(case, 'first_payment_date'))}부터 {_safe(_field(case, 'last_payment_date'))}까지 {_safe(case.get('repayment_months'))}개월간 변제할 계획입니다."),
        _heading("2. 변제에 제공되는 소득 또는 재산"),
        _table(["월 평균 순수입", "인정 생계비", "별도 공제", "월 가용소득", "월 변제예정액", "변제횟수", "명목 총변제예정액"],
               [[_money(calc.get("net_income")), _money(case.get("living_cost")), _money(case.get("deductions")), _money(calc.get("disposable_income")), _money(calc.get("monthly_payment")), case.get("repayment_months"), _money(calc.get("total_payment"))]]),
        _p(f"월 회생위원 보수 및 비용: {_safe(_field(case, 'commissioner_fee_details'))}. 생계비와 공제 항목의 중복 여부를 확인합니다."),
        _heading("3. 재단채권·우선권·별제권 및 미확정채권"),
        _p("해당 채권의 권리 성격·우선순위·담보부족액·유보와 변제방법은 조건부 부속표 및 검증된 배분 결과에 따릅니다." if special else "등록 채권 중 담보·세금·다툼 등 특수분류가 표시된 항목은 없습니다. 담당자는 누락 여부를 최종 확인합니다."),
        _heading("4. 일반 개인회생채권에 대한 변제"),
        _p(f"확정 입력값과 적용 산식 판본 {_safe(calc.get('rule_version'))}에 따른 변제예정 내역은 다음 표와 같습니다."),
        _allocations(case, calc),
        _p(f"총채무 대비 명목 변제율: {_safe(calc.get('repayment_rate'))}%. 원금 기준 또는 현재가치 변제율과 구분합니다."),
        _heading("5. 청산가치와의 비교"),
        _table(["청산가치", "명목 총변제예정액", "현재가치", "비교 기준"], [[_money(calc.get("liquidation_value")), _money(calc.get("total_payment")), _money(calc.get("present_value")), "할인율·기준일이 없는 현재가치는 미산정"]]),
        _heading("6. 변제금 임치 및 지급"),
        _p(f"매월 납입일: {_safe(_field(case, 'payment_day'))}. 회생위원 계좌는 개시결정 시 통지받은 뒤 기재합니다. 현재 신청 단계에서 알 수 없는 계좌를 임의로 생성하지 않습니다."),
        _heading("7. 기타 사항"),
        _p(f"{_safe(_field(case, 'plan_notes'))}"),
    ] + _signature(case))

    secured = [r for r in creditors if r.get("security") == "secured"]
    disputed = [r for r in creditors if r.get("disputed")]
    taxes = [r for r in creditors if r.get("priority") == "tax"]
    if secured:
        add("ANNEX_SECURED", "별제권부 채권 및 이에 준하는 채권의 내역", [
            _table(["계약", "채권자", "원금", "이자", "담보목적물", "담보환가 예상액", "담보상환 예상액", "담보부족 예상액"],
                   [[r.get("contract_id"), r.get("name"), _money(r.get("principal")), _money(r.get("interest")), r.get("collateral"), _money(r.get("collateral_value")), _money(r.get("secured_recovery")), _money(r.get("shortfall"))] for r in secured]),
            _p("담보환가액·담보부족액을 전체 채무 원금에 대한 단순 비례식으로 추정하지 않습니다. 확인된 값과 권리관계 증빙을 사용합니다."),
        ], "D5106 부속서류 1")
    if disputed:
        add("ANNEX_DISPUTED", "다툼이 있거나 예상되는 채권의 내역", [
            _table(["계약", "채권자", "목록상 원금", "목록상 이자", "채권자 주장액", "다툼의 원인", "소송 진행경과", "유보 내역"],
                   [[r.get("contract_id"), r.get("name"), _money(r.get("principal")), _money(r.get("interest")), _money(r.get("claimed_amount")), r.get("dispute_reason", r.get("reason")), r.get("litigation"), r.get("reserve_details")] for r in disputed]),
        ], "D5106 부속서류 2")
    if taxes:
        add("ANNEX_TAX", "조세·공과금 및 우선권 검토 내역", [
            _table(["기관", "계약 / 고지 식별자", "세목", "과세기간", "원금", "가산금 / 이자", "납기", "징수 / 유예", "분류 근거"],
                   [[r.get("name"), r.get("contract_id"), r.get("tax_category"), r.get("tax_period"), _money(r.get("principal")), _money(r.get("interest")), r.get("due_date"), r.get("collection_status"), r.get("classification_basis")] for r in taxes]),
            _p("체납 존재만으로 선납이나 신청 불가를 단정하지 않습니다. 우선권·재단채권·비면책 여부와 적용 시점은 별도 확인합니다."),
        ], "조건부 검토 부속표")
    assignment = [r for r in creditors if r.get("assignment_order")]
    if assignment:
        add("ANNEX_ASSIGNMENT", "전부명령의 내역", [_table(["채권자", "계약", "법원 / 사건", "대상 채권", "제3채무자 송달일", "확정 여부"],
            [[r.get("name"), r.get("contract_id"), r.get("assignment_order"), r.get("assigned_claim"), r.get("service_date"), r.get("order_final")] for r in assignment])], "D5106 부속서류 3")

    add("PAYMENT_SCHEDULE", "개인회생채권 변제예정액표", [
        _p(f"채무자: {_safe(case.get('applicant_name'))} / 변제기간: {_safe(case.get('repayment_months'))}개월"),
        _allocations(case, calc),
        _table(["월 변제예정액", "명목 총변제예정액", "청산가치", "현재가치"], [[_money(calc.get("monthly_payment")), _money(calc.get("total_payment")), _money(calc.get("liquidation_value")), _money(calc.get("present_value"))]]),
    ], "D5110 별지")

    docs = [d for d in case.get("documents", []) if d.get("role") == "evidence"]
    add("EVIDENCE_INDEX", "의뢰인 증빙 및 법원 편철 확인표", [
        _p("의뢰인 원본 수취와 법원 제출용 첨부 완료는 서로 다른 상태입니다. 아래 목록은 원본 증빙 식별용이며 첨부·서명은 실제 제출 전 확인합니다."),
        _table(["자료 ID", "원본 파일", "유형", "대상기간", "원본 페이지", "법원 첨부 상태"],
               [[d.get("id"), d.get("filename", d.get("name")), d.get("document_type"), f"{_safe(d.get('period_start'))} ~ {_safe(d.get('period_end'))}",
                 ", ".join(str(p.get("original_page", p.get("page"))) for p in d.get("pages", []) if p.get("role", d.get("role")) == "evidence") or _safe(d.get("page_count")), "편철 확인 필요"] for d in docs] or [["—", "별도 원본 증빙 등록 필요", UNKNOWN, UNKNOWN, UNKNOWN, "미첨부"]]),
        _p("평가 전용 제출본과 행정자료는 생성 입력·증빙목록에서 제외합니다. 혼합 문서의 증빙 페이지는 원본 페이지 번호로 추적합니다."),
    ], "프로토타입 검토표")

    risk_rows = []
    for risk in assessment.get("risks", []):
        risk_rows.append([risk.get("id"), risk.get("title", risk.get("label")), risk.get("status"),
                          "필수" if risk.get("blocking", risk.get("mandatory", risk.get("required"))) else "경고 / 수동 검토",
                          risk.get("reason", risk.get("description")), risk.get("action", risk.get("next_action"))])
    add("REVIEW_REPORT", "생성본 검토표 및 적용 버전", [
        _table(["항목", "생성 당시 기록"], [["사건 UUID", case.get("id")], ["사건 버전", case.get("revision")],
            ["생성 상태", "미검토 초안"], ["산식 판본", calc.get("rule_version")], ["계산 엔진", calc.get("engine")],
            ["서식 매핑 판본", TEMPLATE_VERSION], ["공식 원본", "templates/originals.json의 URL·SHA-256 참조"],
            ["서식 매핑 검수", "공개 빈 서식 기반 재구성. 사용자 제공 양식과 전문가 최종 대응 검수 별도"],
            ["산식 승인 입력", case.get("rule_approved")], ["서식 승인 입력", case.get("template_approved")],
            ["확인 절차", "실제 생성 파일의 SHA-256별 담당자 확인 → 변호사 최종 확인"],
            ["서명 / 첨부", "생성 완료와 별개로 실제 제출 전 확인"]]),
        _heading("미해결 위험·후속 행동"),
        _table(["식별자", "항목", "상태", "해소 구분", "근거 / 사유", "다음 행동"], risk_rows or [["—", "등록 위험 없음", "자동 완전 검증을 의미하지 않음", "수동 확인", "담당자·변호사 검토", "실제 생성본 대조"]]),
        _heading("남은 확인"),
        *[_p(str(w.get("reason", w.get("title", w))) if isinstance(w, dict) else str(w)) for w in assessment.get("warnings", [])],
        _p("다운로드 후 외부에서 수정한 파일의 내용은 앱에서 추적하지 않습니다. 기존 생성본의 확인을 새로운 파일에 승계하지 않습니다."),
    ], "프로토타입 검토표")
    return models


def render_preview(model):
    parts = [f"<article class=\"document-preview\"><h1>{escape(model['title'])}</h1>"]
    for block in model["blocks"]:
        if block["type"] == "heading":
            parts.append(f"<h2>{escape(block['text'])}</h2>")
        elif block["type"] == "paragraph":
            parts.append(f"<p>{escape(block['text']).replace(chr(10), '<br>')}</p>")
        else:
            parts.append("<table><thead><tr>" + "".join(f"<th>{escape(str(h))}</th>" for h in block["headers"]) + "</tr></thead><tbody>")
            parts.extend("<tr>" + "".join(f"<td>{escape(str(cell))}</td>" for cell in row) + "</tr>" for row in block["rows"])
            parts.append("</tbody></table>")
    return "".join(parts) + "</article>"


def write_docx(model, path):
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal.font.size = Pt(9)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.paragraph_format.space_after = Pt(6)
    title = doc.add_heading(model["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for block in model["blocks"]:
        if block["type"] == "heading":
            doc.add_heading(block["text"], 2)
        elif block["type"] == "paragraph":
            doc.add_paragraph(block["text"])
        else:
            table = doc.add_table(rows=1, cols=len(block["headers"]))
            table.style = "Table Grid"
            for i, heading in enumerate(block["headers"]):
                table.rows[0].cells[i].text = str(heading)
                for run in table.rows[0].cells[i].paragraphs[0].runs:
                    run.bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EEF2")
                table.rows[0].cells[i]._tc.get_or_add_tcPr().append(shade)
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in block["rows"]:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = str(value)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(7.5 if len(block["headers"]) > 6 else 9)
            doc.add_paragraph()
    doc.core_properties.title = model["title"]
    doc.core_properties.subject = f"{model['source_form']} · {TEMPLATE_VERSION}"
    doc.core_properties.author = "Firm Engine prototype retry"
    doc.core_properties.comments = "공개 빈 서식 기반 필드 재구성. 전문가 매핑 검수 별도."
    doc.save(path)


def workbook_model(case, assessment):
    calc = assessment.get("calculations") or {}
    sheets = []
    sheets.append({"name": "계산근거", "rows": [
        ["항목", "확정 결과 / 입력", "산식 / 의미", "출처 / 판본"],
        ["월 총수입", calc.get("monthly_income"), "소득원별 같은 기간의 총수입 합계", "소득원별 원본 참조"],
        ["월 인정경비", calc.get("monthly_expense"), "담당자 인정경비 합계", "소득원별 확인 사유"],
        ["월 순소득", calc.get("net_income"), "월 총수입 - 월 인정경비", calc.get("rule_version")],
        ["인정 생계비", case.get("living_cost"), "사건별 확정 입력", _field(case,"living_cost_reason")],
        ["별도 공제", case.get("deductions"), "사건별 확정 입력", _field(case,"deductions_reason")],
        ["월 가용소득", calc.get("disposable_income"), "월 순소득 - 인정 생계비 - 공제", calc.get("rule_version")],
        ["재산 청산가치", calc.get("liquidation_value"), "재산별 순가치 - 확인된 면제 인정액", "재산 목록 / 평가 출처"],
        ["채무 원금 합계", calc.get("total_principal"), "계약별 원금 합계", "채권자목록"],
        ["채무 이자 합계", calc.get("total_interest"), "계약별 이자 합계", "채권자목록"],
        ["총채무", calc.get("total_debt"), "원금 합계 + 이자 합계", "채권자목록"],
        ["변제기간(개월)", case.get("repayment_months"), "사건별 승인 기간", "고정 36개월 아님"],
        ["월 변제예정액", calc.get("monthly_payment"), "검증된 계산 결과", calc.get("engine")],
        ["명목 총변제예정액", calc.get("total_payment"), "월별 배분 결과의 합계", calc.get("rule_version")],
        ["명목 변제율(%)", calc.get("repayment_rate"), "명목 총변제예정액 / 총채무 × 100", "현재가치 변제율과 구분"],
        ["현재가치", calc.get("present_value"), "비교 기준일·할인율 미확정 시 미산정", "확인 필요"],
        ["사건 UUID", case.get("id"), "사건 식별자", "영속 사건 원장"],
        ["사건 버전", case.get("revision"), "생성 시 입력 버전", "원본 변경 시 재검수"],
        ["서식 매핑 판본", TEMPLATE_VERSION, "공개 빈 양식 기반 재구성", "전문가 최종 대응 검수 별도"],
    ]})
    income_rows = [["소득원 ID", "사업체 / 직장", "유형", "대상 시작일", "대상 종료일", "월 총수입", "월 인정경비", "월 순소득", "원본 자료 ID", "원본 페이지", "원문 근거", "확인 상태"]]
    for row in case.get("incomes", []):
        gross, expense = _number(row.get("monthly_gross")), _number(row.get("monthly_expense"))
        income_rows.append([row.get("id"), row.get("label"), row.get("kind"), row.get("period_start"), row.get("period_end"), gross, expense,
                            gross-expense if gross is not None and expense is not None else None,
                            row.get("source_document_id"), row.get("source_page"), row.get("source_quote"), "확인" if row.get("confirmed") else "미확인"])
    sheets.append({"name": "소득원별근거", "rows": income_rows})
    monthly_rows = [["귀속월", "소득원 ID", "소득원", "월 총수입", "월 인정경비", "월 순소득", "산정 범위"]]
    for row in case.get("incomes", []):
        start, end = str(row.get("period_start") or "")[:7], str(row.get("period_end") or "")[:7]
        try:
            sy, sm = (int(n) for n in start.split("-")); ey, em = (int(n) for n in end.split("-"))
            if not (1 <= sm <= 12 and 1 <= em <= 12): raise ValueError()
            count = (ey-sy)*12 + em-sm+1
            if not 1 <= count <= 240: raise ValueError()
            gross, expense = _number(row.get("monthly_gross")), _number(row.get("monthly_expense"))
            for offset in range(count):
                value = sy*12+sm-1+offset
                month = f"{value//12:04d}-{value%12+1:02d}"
                monthly_rows.append([month, row.get("id"), row.get("label"), gross, expense, gross-expense if gross is not None and expense is not None else None,
                                     "확정 월 평균값 배치; 개별 월 실제 추출값과 구분"])
        except (ValueError, TypeError):
            monthly_rows.append([UNKNOWN, row.get("id"), row.get("label"), None, None, None, "대상기간 확인 필요"])
    sheets.append({"name": "월별소득", "rows": monthly_rows})
    allocations = _allocations(case, calc)
    sheets.append({"name": "채권별배분", "rows": [allocations["headers"]] + allocations["rows"]})
    fact_rows = [["사실 ID", "항목", "값", "단위", "대상 객체", "자료 ID", "원본 페이지", "원문 발췌", "확인 상태", "확인 사유"]]
    for row in case.get("facts", []):
        fact_rows.append([row.get("id"), row.get("label", row.get("key")), row.get("value"), row.get("unit"), row.get("object_id"), row.get("document_id"), row.get("page"), row.get("quote"), row.get("status"), row.get("reason")])
    sheets.append({"name": "원본추적", "rows": fact_rows})
    risk_rows = [["위험 ID", "항목", "현재 상태", "필수 해소", "근거 / 사유", "다음 행동"]]
    for row in assessment.get("risks", []):
        risk_rows.append([row.get("id"), row.get("title", row.get("label")), row.get("status"), row.get("blocking", row.get("mandatory")), row.get("reason", row.get("description")), row.get("action", row.get("next_action"))])
    sheets.append({"name": "검토정보", "rows": risk_rows})
    if case.get("diagnostic_mode"):
        for sheet in sheets:
            sheet["banner"] = DIAGNOSTIC_BANNER
    return sheets


def write_workbook(sheets, path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for model in sheets:
        sheet = workbook.create_sheet(model["name"][:31])
        if model.get("banner"):
            sheet.append([model["banner"]])
        for row in model["rows"]:
            safe = []
            for value in row:
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False)
                if value is None or value == "": value = UNKNOWN
                # User-entered evidence is text, never an executable spreadsheet formula.
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                    value = "'" + value
                safe.append(value)
            sheet.append(safe)
        header_index = 2 if model.get("banner") else 1
        sheet.freeze_panes = f"A{header_index+1}"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[header_index]:
            cell.fill = PatternFill("solid", fgColor="1B3446")
            cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
        if model.get("banner"):
            sheet.cell(1,1).font = Font(name="맑은 고딕", bold=True, color="9C2525")
        for row in sheet.iter_rows(min_row=header_index+1):
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.##"
        for index in range(1, sheet.max_column+1):
            sheet.column_dimensions[get_column_letter(index)].width = 24 if index < 6 else 32
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    workbook.save(path)


def generate_package(case: dict, assessment: dict, output_dir: Path) -> list[dict]:
    """Write a new immutable generation directory; parent enforces readiness gates."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    if any(output_dir.iterdir()):
        raise ValueError("생성물 덮어쓰기 금지: 비어 있는 새 버전 디렉터리가 필요합니다.")
    models = content_models(case, assessment)
    artifacts = []
    renderer_sha256 = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    for model in models:
        path = output_dir / f"{model['code']}.docx"
        write_docx(model, path)
        artifacts.append({"title": model["title"], "kind": "docx", "path": str(path),
                          "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "preview_html": render_preview(model),
                          "template_version": TEMPLATE_VERSION, "code": model["code"]})
    sheets = workbook_model(case, assessment)
    workbook_path = output_dir / "CALCULATION_BASIS.xlsx"
    write_workbook(sheets, workbook_path)
    workbook_preview = "".join(render_preview({"title": sheet["name"], "blocks": ([_p(sheet["banner"])] if sheet.get("banner") else []) + [_table(sheet["rows"][0], sheet["rows"][1:])]}) for sheet in sheets)
    artifacts.append({"title": "계산근거표 및 원본 추적", "kind": "xlsx", "path": str(workbook_path),
                      "sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(), "preview_html": workbook_preview,
                      "template_version": TEMPLATE_VERSION, "code": "CALCULATION_BASIS"})
    for artifact in artifacts:
        artifact.update({"renderer_version": RENDERER_VERSION, "renderer_sha256": renderer_sha256})
    originals = ROOT / "templates" / "originals.json"
    source_manifest = json.loads(originals.read_text(encoding="utf-8")) if originals.exists() else []
    manifest = {"case_id": case.get("id"), "case_revision": case.get("revision"),
                "generated_at": datetime.now(timezone.utc).isoformat(), "review_status_at_generation": "미검토 초안",
                "diagnostic_mode": bool(case.get("diagnostic_mode")),
                "generation_result": "diagnostic_candidate_render_only" if case.get("diagnostic_mode") else "court_draft_rendered",
                "renderer_version": RENDERER_VERSION, "renderer_sha256": renderer_sha256,
                "template_version": TEMPLATE_VERSION, "rule_version": assessment.get("calculations", {}).get("rule_version"),
                "sources": source_manifest,
                "artifacts": [{key:value for key,value in item.items() if key not in ("preview_html", "path")} | {"filename": Path(item["path"]).name} for item in artifacts]}
    (output_dir / "generation_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    # Retained content model enables an independent evaluator to compare exact document bodies.
    (output_dir / "content_model.json").write_text(json.dumps(models, ensure_ascii=False, indent=2), encoding="utf-8")
    zip_path = output_dir / "COURT_DRAFT_PACKAGE.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(output_dir.iterdir()):
            if path != zip_path:
                archive.write(path, path.name)
    artifacts.append({"title": "전체 초안 패키지", "kind": "zip", "path": str(zip_path),
                      "sha256": hashlib.sha256(zip_path.read_bytes()).hexdigest(),
                      "preview_html": "<p>개별 DOCX·XLSX와 생성 버전·원본 해시 매니페스트를 포함합니다. 생성 당시 미검토 초안입니다.</p>",
                      "template_version": TEMPLATE_VERSION, "code": "COURT_DRAFT_PACKAGE",
                      "renderer_version": RENDERER_VERSION, "renderer_sha256": renderer_sha256})
    return artifacts
